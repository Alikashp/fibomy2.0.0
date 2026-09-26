"""
Спайк шрифтов (ТЗ, раздел 9, фаза 0; ТЗ 6.3).

Строит out/fonts_matrix.pptx: по слайду на каждый шрифт-кандидат, на слайде —
тестовая строка на каждом из 15 языков (ТЗ 11, вопрос 3) и строка символов
(валюты СНГ, типографика). Плюс слайд с подстановкой для армянского и грузинского.

Затем проверяет серверную сторону — то, что увидит пользователь в PDF:
  1. fc-match: каким файлом шрифта LibreOffice на сервере заменит кандидата
     и совместим ли он по метрикам (иначе переносы строк в PDF и в PowerPoint
     разойдутся);
  2. покрытие глифов (cmap) этого файла для каждого языка — каких символов нет;
  3. какими шрифтами LibreOffice реально набрал каждую строку в PDF
     (если не основным — значит, сработала подстановка по глифам).
Результаты: out/fonts_report.json и out/fonts_coverage.md.

Как шрифт выглядит в PowerPoint, Keynote и Google Slides, отсюда проверить
нельзя — это ручной чек-лист в docs/SPIKE_PPTX.md.

Запуск из корня репозитория:
    python spikes/fonts_spike.py
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import OUT, pdf_to_pngs, pptx_to_pdf, rgb, u  # noqa: E402

# Язык → (тег для a:rPr lang, тестовая строка). Строки собраны так, чтобы
# в каждой были буквы, специфичные для языка.
LANGS: dict[str, tuple[str, str]] = {
    "ru": ("ru-RU", "Съешь же ещё этих мягких французских булок, да выпей чаю. Ё ё"),
    "en": ("en-US", "The quick brown fox jumps over the lazy dog. “Quotes” — 1st, 2nd"),
    "uk": ("uk-UA", "Україна: Ґ ґ Є є І і Ї ї, м’ята, під’їзд"),
    "be": ("be-BY", "Рэспубліка Беларусь: Ў ў І і Ё ё, вераб’я"),
    "kk": ("kk-KZ", "Қазақстан: Ә ә Ғ ғ Қ қ Ң ң Ө ө Ұ ұ Ү ү Һ һ І і"),
    "uz": ("uz-Latn-UZ", "Oʻzbekiston: Oʻ oʻ Gʻ gʻ (U+02BB) · O‘ o‘ G‘ g‘ (U+2018) · Sh ch"),
    "ky": ("ky-KG", "Кыргызстан: Ң ң Ө ө Ү ү — Бишкек, Ысык-Көл"),
    "tg": ("tg-Cyrl-TJ", "Тоҷикистон: Ғ ғ Ӣ ӣ Қ қ Ӯ ӯ Ҳ ҳ Ҷ ҷ — Душанбе"),
    "az": ("az-Latn-AZ", "Azərbaycan: Ə ə Ğ ğ I ı İ i Ö ö Ş ş Ü ü Ç ç"),
    "hy": ("hy-AM", "Հայաստան — Երևան, Ա ա Բ բ Գ գ Ֆ ֆ և"),
    "ka": ("ka-GE", "საქართველო — თბილისი, ქართული ანბანი"),
    "es": ("es-ES", "El pingüino Wenceslao hizo kilómetros. ¿Ñ ñ? ¡Á é í ó ú!"),
    "de": ("de-DE", "Zwölf Boxkämpfer jagen Viktor über den Deich. Ä Ö Ü ß ẞ „Zitat“"),
    "fr": ("fr-FR", "Voix ambiguë d’un cœur qui préfère les kiwis. Œ œ Ç ç « »"),
    "tr": ("tr-TR", "Pijamalı hasta yağız şoföre çabucak güvendi. İ ı Ğ ğ Ş ş"),
}
SYMBOLS = ("sym", "0123456789 % ‰ № € $ £ ₽ ₸ ₼ ₾ ֏ ₺ — – … • × ± ≈ ≤ ≥ → ↑")

# Кандидаты: есть в Windows и macOS с Office и в меню шрифтов Google Slides
# (Segoe UI — только Windows, для сравнения).
CANDIDATES = [
    "Arial", "Calibri", "Segoe UI", "Verdana", "Tahoma",
    "Trebuchet MS", "Georgia", "Times New Roman",
]
# Метрически совместимые свободные клоны — их можно положить в образ воркера,
# и тогда PDF из LibreOffice переносит строки так же, как PowerPoint.
METRIC_CLONES = {
    "Arial": "Liberation Sans",
    "Times New Roman": "Liberation Serif",
    "Calibri": "Carlito",
    "Cambria": "Caladea",
    "Georgia": "Gelasio",
    "Segoe UI": "Selawik",
}
# Для армянского и грузинского (ТЗ 6.3: допускается шрифт-подстановка по языку)
SCRIPT_FALLBACKS = {
    "hy": ["Sylfaen", "Arial", "Noto Sans Armenian"],
    "ka": ["Sylfaen", "Arial", "Noto Sans Georgian"],
}


def _add_line(slide, y, label, text, font, lang_tag, size=16):
    lbl = slide.shapes.add_textbox(u(80), u(y), u(90), u(40))
    _fill(lbl, label, "Arial", "en-US", size, "5B6570")
    box = slide.shapes.add_textbox(u(180), u(y), u(1660), u(40))
    _fill(box, text, font, lang_tag, size, "1B1F24")


def _fill(shape, text, font, lang_tag, size, color):
    tf = shape.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", lang_tag)
    # Имя шрифта и для «сложных» письменностей, и для восточноазиатских —
    # иначе PowerPoint может взять шрифт темы для части символов.
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", font)


def build_matrix() -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12_192_000), Emu(6_858_000)
    for font in CANDIDATES:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        title = s.shapes.add_textbox(u(80), u(40), u(1760), u(60))
        _fill(title, f"{font} — 15 языков + символы", font, "ru-RU", 24, "1F4E79")
        title.text_frame.paragraphs[0].runs[0].font.bold = True
        y = 120
        for code, (tag, text) in LANGS.items():
            _add_line(s, y, code, text, font, tag)
            y += 54
        _add_line(s, y, SYMBOLS[0], SYMBOLS[1], font, "ru-RU")
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title = s.shapes.add_textbox(u(80), u(40), u(1760), u(60))
    _fill(title, "Армянский и грузинский: шрифт-подстановка по языку", "Arial", "ru-RU", 24, "1F4E79")
    y = 140
    for code, fonts in SCRIPT_FALLBACKS.items():
        tag, text = LANGS[code]
        for font in fonts:
            _add_line(s, y, code, f"{text}   [{font}]", font, tag, size=20)
            y += 70
    path = OUT / "fonts_matrix.pptx"
    prs.save(path)
    return path


# ── Серверная сторона ───────────────────────────────────────────────────────

def fc_match(family: str) -> tuple[str, str]:
    out = subprocess.run(["fc-match", "-f", "%{family[0]}|%{file}", family],
                         capture_output=True, text=True, check=True).stdout
    name, path = out.split("|", 1)
    return name, path


def missing_glyphs(font_path: str, text: str) -> list[str]:
    from fontTools.ttLib import TTFont

    font = TTFont(font_path, fontNumber=0, lazy=True)
    cmap = font.getBestCmap() or {}
    return sorted({ch for ch in text if not ch.isspace() and ord(ch) not in cmap})


def pdf_line_fonts(pdf: Path) -> dict[int, dict[str, list[str]]]:
    """Страница → {первые символы строки: шрифты, которыми она набрана}."""
    import pymupdf

    result = {}
    with pymupdf.open(pdf) as doc:
        for i, page in enumerate(doc, start=1):
            lines = {}
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    text = "".join(s["text"] for s in line["spans"]).strip()
                    if text:
                        lines[text[:24]] = sorted({s["font"].split("+")[-1] for s in line["spans"]})
            result[i] = lines
    return result


def main():
    OUT.mkdir(exist_ok=True)
    pptx_path = build_matrix()
    pdf_path, elapsed = pptx_to_pdf(pptx_path, OUT)
    pdf_to_pngs(pdf_path, OUT / "previews", width_px=1600)
    line_fonts = pdf_line_fonts(pdf_path)

    rows = []
    for page, font in enumerate(CANDIDATES, start=1):
        server_name, server_file = fc_match(font)
        per_lang = {}
        samples = [(code, text) for code, (_, text) in LANGS.items()] + [SYMBOLS]
        for code, text in samples:
            miss = missing_glyphs(server_file, text)
            used = next((fonts for key, fonts in line_fonts.get(page, {}).items() if text.startswith(key[:12])), [])
            per_lang[code] = {"missing": miss, "pdf_fonts": used}
        rows.append({
            "font": font,
            "server_substitute": server_name,
            "metric_compatible": METRIC_CLONES.get(font) == server_name,
            "known_metric_clone": METRIC_CLONES.get(font),
            "langs": per_lang,
        })

    report = {"convert_sec": round(elapsed, 2), "fonts": rows}
    (OUT / "fonts_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "fonts_coverage.md").write_text(render_md(rows), encoding="utf-8")
    print(render_md(rows))


def render_md(rows) -> str:
    codes = list(LANGS) + ["sym"]
    lines = [
        "# Покрытие глифов на сервере (LibreOffice → PDF)",
        "",
        "Сгенерировано `spikes/fonts_spike.py`. «✓» — у шрифта-замены на сервере есть все символы",
        "тестовой строки; иначе — список отсутствующих символов (LibreOffice доберёт их из другого",
        "шрифта, «квадратиков» не будет, но строка окажется набрана смесью шрифтов).",
        "",
        "| Кандидат | Замена на сервере | Метрики совпадают | " + " | ".join(codes) + " |",
        "|---|---|---|" + "---|" * len(codes),
    ]
    for r in rows:
        cells = []
        for c in codes:
            miss = r["langs"][c]["missing"]
            cells.append("✓" if not miss else " ".join(miss))
        metric = "да" if r["metric_compatible"] else (f"нет (клон: {r['known_metric_clone']})" if r["known_metric_clone"] else "нет")
        lines.append(f"| {r['font']} | {r['server_substitute']} | {metric} | " + " | ".join(cells) + " |")
    return "\n".join(lines) + "\n"


if __name__ == "__main__":
    main()
