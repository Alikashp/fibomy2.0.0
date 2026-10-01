"""INGEST: материал пользователя → SourceDigest (03_ARCHITECTURE.md, 3; ТЗ 3.1).

- текст сообщения, .txt (chardet), .docx (абзацы, заголовки, таблицы), .pdf
  (текст и таблицы через pdfplumber.find_tables, до 50 страниц), .pptx (текст,
  заметки, таблицы, данные нативных диаграмм);
- таблицы и ряды чисел в тексте → datasets с адресами ячеек;
- docx и pptx — zip: до разбора проверяются размер после распаковки и число файлов;
- разбор — в потоке с таймаутом 15 с;
- текст длиннее 40 000 знаков — начало документа, truncated и предупреждение;
- PDF без текстового слоя — ошибка SCAN_WITHOUT_TEXT.
"""

import asyncio
import io
import logging
import re
import zipfile

from core.ingest.tables import build_dataset, series_dataset
from core.models.digest import Fragment, Source, SourceDigest
from core.models.request import DeckRequest

logger = logging.getLogger(__name__)

MAX_CHARS = 40_000
MAX_PDF_PAGES = 50
MAX_PPTX_SLIDES = 100
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024
MAX_ZIP_FILES = 5000
PARSE_TIMEOUT = 15.0
MIN_TEXT_CHARS = 40          # меньше — это скан или пустой файл

KIND_BY_MIME = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "text/plain": "txt",
}


class IngestError(RuntimeError):
    """Материал не прочитан; code — SCAN_WITHOUT_TEXT | BAD_FILE | PARSE_TIMEOUT | FILE_UNSUPPORTED."""

    def __init__(self, code: str, message: str = ""):
        super().__init__(message or code)
        self.code = code


class _Builder:
    """Собирает фрагменты и наборы данных с последовательными id."""

    def __init__(self):
        self.fragments: list[Fragment] = []
        self.datasets = []
        self.pending_title: str | None = None

    def text(self, text: str, kind: str = "paragraph", **loc) -> None:
        text = " ".join(text.split())
        if not text:
            return
        self.fragments.append(Fragment(id=f"f{len(self.fragments) + 1}", text=text, kind=kind, loc=loc))
        ds = series_dataset(f"ds{len(self.datasets) + 1}", text, loc={"fragment": self.fragments[-1].id, **loc})
        if ds is not None:
            self.datasets.append(ds)
        # «Таблица 2. Выручка по каналам» перед таблицей — её подпись
        self.pending_title = text if (kind == "heading" or re.match(r"^(таблица|table|кесте)\s*\d", text, re.I)
                                      or len(text) < 90) else None

    def table(self, rows: list[list[str]], origin: str, **loc) -> None:
        ds = build_dataset(f"ds{len(self.datasets) + 1}", rows, origin, title=self.pending_title, loc=loc)
        if ds is not None:
            self.datasets.append(ds)
        else:  # таблица без чисел — текстом
            for r in rows:
                line = " | ".join(" ".join((c or "").split()) for c in r if (c or "").strip())
                if line:
                    self.text(line, "list_item", **loc)
        self.pending_title = None


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
    except zipfile.BadZipFile as e:
        raise IngestError("BAD_FILE", "not a zip container") from e
    if len(infos) > MAX_ZIP_FILES or sum(i.file_size for i in infos) > MAX_UNZIPPED_BYTES:
        raise IngestError("BAD_FILE", "zip too large after decompression")


def _docx(data: bytes, b: _Builder) -> dict:
    from docx import Document
    from docx.table import Table

    _check_zip(data)
    doc = Document(io.BytesIO(data))
    for n, block in enumerate(doc.iter_inner_content()):
        if isinstance(block, Table):
            rows, prev = [], None
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    if cell._tc is prev:      # объединённая ячейка повторяется — схлопываем
                        continue
                    prev = cell._tc
                    cells.append(cell.text)
                rows.append(cells)
            b.table(rows, "docx_table", para=n)
        else:
            style = (block.style.name if block.style is not None else "") or ""
            kind = "heading" if style.lower().startswith(("heading", "title", "заголов")) else \
                "list_item" if "list" in style.lower() else "paragraph"
            b.text(block.text, kind, para=n)
    return {}


def _merge_lines(text: str) -> list[str]:
    """Строки PDF → абзацы: строка без точки в конце продолжается следующей со строчной буквы."""
    paras, cur = [], ""
    for line in (text or "").split("\n"):
        line = line.strip()
        if not line:
            if cur:
                paras.append(cur)
                cur = ""
            continue
        if cur and not re.search(r"[.!?:;]$", cur) and (line[0].islower() or line[0] in ",)–—-"):
            cur += " " + line
        else:
            if cur:
                paras.append(cur)
            cur = line
    if cur:
        paras.append(cur)
    return paras


def _pdf(data: bytes, b: _Builder) -> dict:
    import pdfplumber

    def inside(obj, bboxes):
        cx, cy = (obj["x0"] + obj["x1"]) / 2, (obj["top"] + obj["bottom"]) / 2
        return any(x0 <= cx <= x1 and t <= cy <= bt for x0, t, x1, bt in bboxes)

    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as e:
        raise IngestError("BAD_FILE", f"pdf: {e}") from e
    warnings = []
    with pdf:
        pages = pdf.pages
        if len(pages) > MAX_PDF_PAGES:
            warnings.append(f"pdf_pages:{MAX_PDF_PAGES}/{len(pages)}")
        for pno, page in enumerate(pages[:MAX_PDF_PAGES], start=1):
            tables = sorted(page.find_tables(), key=lambda t: t.bbox[1])
            if not tables:
                for para in _merge_lines(page.extract_text() or ""):
                    b.text(para, page=pno)
                continue
            bboxes = [t.bbox for t in tables]
            text_page = page.filter(lambda o: o.get("object_type") != "char" or not inside(o, bboxes))
            x0, top0, x1, bottom0 = page.bbox
            cursor = top0
            for t in tables:
                _, t_top, _, t_bottom = t.bbox
                if t_top > cursor:
                    for para in _merge_lines(text_page.crop((x0, cursor, x1, t_top)).extract_text() or ""):
                        b.text(para, page=pno)
                b.table([[c or "" for c in row] for row in t.extract()], "pdf_table", page=pno)
                cursor = max(cursor, t_bottom)
            if cursor < bottom0:
                for para in _merge_lines(text_page.crop((x0, cursor, x1, bottom0)).extract_text() or ""):
                    b.text(para, page=pno)
        return {"pages": len(pages), "warnings": warnings}


def _pptx(data: bytes, b: _Builder) -> dict:
    from pptx import Presentation

    _check_zip(data)
    prs = Presentation(io.BytesIO(data))
    warnings = []
    slides = list(prs.slides)
    if len(slides) > MAX_PPTX_SLIDES:
        warnings.append(f"pptx_slides:{MAX_PPTX_SLIDES}/{len(slides)}")
    for sno, slide in enumerate(slides[:MAX_PPTX_SLIDES], start=1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    b.text("".join(r.text for r in p.runs) or p.text if hasattr(p, "text") else "", slide=sno)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                b.table(rows, "pptx_table", slide=sno)
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                chart = shape.chart
                try:
                    plot = chart.plots[0]
                    cats = [str(c) for c in plot.categories]
                    series = list(plot.series)
                    header = [""] + [s.name or f"Серия {i + 1}" for i, s in enumerate(series)]
                    rows = [header] + [[cat] + [_fmt(s.values[i]) for s in series] for i, cat in enumerate(cats)]
                    b.pending_title = chart.chart_title.text_frame.text if chart.has_title else b.pending_title
                    b.table(rows, "pptx_chart", slide=sno)
                except Exception as e:  # диаграмма без данных или нестандартная
                    warnings.append(f"pptx_chart_unreadable:slide{sno}")
                    logger.info("Chart data not readable", extra={"slide": sno, "error": str(e)})
        if slide.has_notes_slide:
            for line in slide.notes_slide.notes_text_frame.text.split("\n"):
                b.text(line, "note", slide=sno)
    return {"pages": len(slides), "warnings": warnings}


def _fmt(value) -> str:
    if value is None:
        return ""
    return str(int(value)) if float(value) == int(value) else str(value).replace(".", ",")


def _plain(text: str, b: _Builder) -> dict:
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        lines = [x.strip() for x in block.split("\n") if x.strip()]
        for line in lines:
            kind = "list_item" if re.match(r"^([-•*–—]|\d+[.)])\s", line) else "paragraph"
            b.text(re.sub(r"^[-•*–—]\s+", "", line), kind)
    return {}


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        import chardet
        enc = chardet.detect(data).get("encoding") or "utf-8"
        return data.decode(enc, errors="replace")


def parse(data: bytes, kind: str, name: str | None = None) -> SourceDigest:
    """Синхронный разбор (вызывается в потоке). kind: text | txt | docx | pdf | pptx."""
    b = _Builder()
    try:
        if kind in ("text", "txt"):
            meta = _plain(_decode(data), b)
        elif kind == "docx":
            meta = _docx(data, b)
        elif kind == "pdf":
            meta = _pdf(data, b)
        elif kind == "pptx":
            meta = _pptx(data, b)
        else:
            raise IngestError("FILE_UNSUPPORTED", kind)
    except IngestError:
        raise
    except Exception as e:
        raise IngestError("BAD_FILE", f"{kind}: {type(e).__name__}: {e}") from e

    total = sum(len(f.text) for f in b.fragments)
    table_chars = sum(len(c.raw) for ds in b.datasets for r in ds.rows for c in r.cells.values())
    if total + table_chars < MIN_TEXT_CHARS:
        if kind == "pdf":
            raise IngestError("SCAN_WITHOUT_TEXT", "no text layer")
        raise IngestError("BAD_FILE", "no text")

    fragments, used, truncated = [], 0, False
    for f in b.fragments:
        if used + len(f.text) > MAX_CHARS:
            truncated = True
            break
        fragments.append(f)
        used += len(f.text)
    warnings = list(meta.get("warnings", []))
    if truncated:
        warnings.append(f"truncated:{used}/{total}")
        kept = {f.id for f in fragments}
        datasets = [d for d in b.datasets if d.loc.get("fragment", next(iter(kept), "")) in kept or "fragment" not in d.loc]
    else:
        datasets = b.datasets
    return SourceDigest(
        mode="material", topic="",
        source=Source(kind="text" if kind == "text" else kind, name=name, pages=meta.get("pages"),
                      chars_total=total, chars_used=used, truncated=truncated),
        fragments=fragments, datasets=datasets, warnings=warnings,
    )


async def ingest(request: DeckRequest, data: bytes | None = None) -> SourceDigest:
    """DeckRequest (+ байты материала) → SourceDigest. По теме — пустой digest с темой."""
    if request.mode == "topic":
        return SourceDigest.for_topic(request.input.topic)
    material = request.input.material
    if data is None:
        raise IngestError("BAD_FILE", "material bytes are missing")
    kind = "text" if material.kind == "text" else KIND_BY_MIME.get(material.mime or "", "")
    if not kind:
        ext = (material.name or "").rsplit(".", 1)[-1].lower()
        kind = ext if ext in ("pdf", "docx", "pptx", "txt") else ""
    if not kind:
        raise IngestError("FILE_UNSUPPORTED", material.mime or "unknown")
    try:
        digest = await asyncio.wait_for(asyncio.to_thread(parse, data, kind, material.name), timeout=PARSE_TIMEOUT)
    except asyncio.TimeoutError as e:
        raise IngestError("PARSE_TIMEOUT", f"parse > {PARSE_TIMEOUT:.0f}s") from e
    digest.topic = request.input.topic
    logger.info("Ingested", extra={"kind": kind, "fragments": len(digest.fragments), "datasets": len(digest.datasets),
                                   "chars": digest.source.chars_used, "truncated": digest.source.truncated})
    return digest
