"""Заполненность зоны контента (замечание владельца, сессия 5: «пустая нижняя половина слайдов»; D-073).

После подбора кегля блок контента стоит по центру зоны (compact: center в YAML). Если он
занимает меньше FILL_MIN (40%) зоны, карточки, колонки и панели растягиваются по высоте —
шагами до всей области варианта, пока доля не дойдёт до 40%. Списки без подложек
в несколько строк растягиваются так же: строки расходятся по зоне, текст каждой — у верха своей
строки; колонки без подложек в один ряд остаются блоком по центру зоны.

Доля считается по самому слайду, нарисованному рендерером (та же мера, что у галереи и
финальной проверки, core.checks.layout.fill_share).
"""

from dataclasses import replace

from pptx import Presentation

from core.checks.layout import FILL_EXEMPT, FILL_KINDS, FILL_MIN, fill_share
from core.models import bind as B
from core.models.deck import DeckSpec, Slide
from core.models.layout import LayoutSpec
from core.models.theme import STYLE_STEPS, Theme

STEPS = 6


def slide_fill(slide: Slide, spec: DeckSpec, layout: LayoutSpec) -> float:
    """Доля зоны контента у слайда в текущем виде (рисует один слайд в память)."""
    from core.render.pptx import renderer as R
    from core.models.theme import load_theme
    prs = Presentation()
    prs.slide_width, prs.slide_height = R.Emu(R.SLIDE_W_EMU), R.Emu(R.SLIDE_H_EMU)
    ctx = R._Ctx(spec, load_theme(spec.meta.theme_id), {})
    R._render_slide(prs, ctx, slide, layout, max(len(spec.slides), slide.index), False)
    return fill_share(prs.slides[0])


def stretchable(layout: LayoutSpec, n: int) -> bool:
    """Растягивать есть смысл карточки (подложка на всю ячейку) и строки списка (строки расходятся
    по зоне). Колонки без подложек в один ряд остаются блоком по центру: растянутые, они
    вернули бы текст к верху зоны с пустым низом."""
    arr = layout.arrangements[n]
    return int(arr["grid"].get("rows", 1)) > 1 or any(e.get("type") == "card" for e in arr["elements"])


def grow_styles(slide: Slide, spec: DeckSpec, layout: LayoutSpec, theme: Theme) -> bool:
    """Мало текста: тексты пунктов на шаг крупнее базового стиля, если так влезают все. → изменилось ли."""
    from core.fitting.fitter import measure
    items = B.items_of(layout.items_bind, slide)
    if not items or len(items) not in layout.arrangements:
        return False
    per_item = layout.item_elements(len(items))
    changed = False
    for pos, template in enumerate(per_item[0]):
        if template.type != "text" or not template.bind or template.bind.startswith(("$", "const:")):
            continue
        key = f"item.{template.name}"
        current = slide.fit.styles.get(key)
        if current != template.style or STYLE_STEPS.index(current) == 0 or template.suffix:
            continue
        bigger = STYLE_STEPS[STYLE_STEPS.index(current) - 1]
        if bigger in ("hero", "display"):
            continue
        texts = [(per_item[i][pos], B.text_of(per_item[i][pos], slide, spec, item=it, index=i))
                 for i, it in enumerate(items)]
        if all(not t or measure(replace(e, max_lines=e.max_lines), t, bigger, theme).fits
               for e, t in texts):
            slide.fit.styles[key] = bigger
            changed = True
    return changed


def applies(slide: Slide, layout: LayoutSpec) -> bool:
    return layout.kind in FILL_KINDS and layout.id not in FILL_EXEMPT


def ensure_fill(slide: Slide, spec: DeckSpec, layout: LayoutSpec, theme: Theme) -> float | None:
    """Растягивает блок контента до нижней границы заполненности; мало и этого — тексты пунктов на шаг
    крупнее. → итоговая доля или None."""
    if not applies(slide, layout):
        return None
    share = slide_fill(slide, spec, layout)
    if share >= FILL_MIN:
        return share
    share = _stretch(slide, spec, layout, share)
    if share < FILL_MIN and grow_styles(slide, spec, layout, theme):
        from core.fitting.fitter import compact_items
        compact_items(slide, spec, layout, theme)
        share = slide_fill(slide, spec, layout)
        if share < FILL_MIN:
            share = _stretch(slide, spec, layout, share)
    return share


def _stretch(slide: Slide, spec: DeckSpec, layout: LayoutSpec, share: float) -> float:
    items = B.items_of(layout.items_bind, slide)
    if items and len(items) in layout.arrangements and slide.fit.items_area:
        grid = layout.arrangements[len(items)]["grid"]
        if not stretchable(layout, len(items)):
            return share
        ax, ay, aw, ah = (float(v) for v in grid["area"])
        _, top, _, height = slide.fit.items_area
        center = grid.get("compact") == "center"
        for k in range(1, STEPS + 1):
            h = height + (ah - height) * k / STEPS
            t = ay + (ah - h) / 2 if center else ay
            slide.fit.items_area = [ax, round(t, 1), aw, round(h, 1)]
            share = slide_fill(slide, spec, layout)
            if share >= FILL_MIN:
                break
        if slide.fit.items_area[3] >= ah - 0.5:
            slide.fit.items_area = None                     # вся область варианта — как в YAML
        return share
    if layout.compact and slide.fit.boxes:
        boxes = dict(slide.fit.boxes)
        els = {e.name: e for e in layout.elements()}
        panels = layout.compact.get("panels", [])
        for k in range(1, STEPS + 1):
            for name in panels:
                e = els[name]
                x, y, w, h = slide.fit.boxes[name]
                full = e.box.h
                grown = h + (full - h) * k / STEPS
                dy = (e.box.y + (full - grown) / 2) - y if layout.compact.get("align", "center") == "center" else 0
                boxes[name] = [x, round(y + dy, 1), w, round(grown, 1)]
            # пункты и шапки — вместе с панелями (по центру зоны)
            shift = boxes[panels[0]][1] - slide.fit.boxes[panels[0]][1] if panels else 0
            for name, b in slide.fit.boxes.items():
                if name not in panels:
                    boxes[name] = [b[0], round(b[1] + shift, 1), b[2], b[3]]
            trial = slide.model_copy(deep=True)
            trial.fit.boxes = boxes
            share = slide_fill(trial, spec, layout)
            if share >= FILL_MIN:
                break
        slide.fit.boxes = boxes
        return share
    return share
