"""Финальная проверка вёрстки по готовому PPTX (ТЗ 7.1; сессия 5, D-072).

Проверяется сам файл, который получит пользователь, — каждая текстовая рамка с её кеглем
и шириной по метрикам реального шрифта (Liberation Sans = Arial, D-027), а не DeckSpec:
- overflow — строк больше, чем вмещает высота рамки (текст вылез за рамку);
- word — слово не помещается в ширину рамки («Непротиворечивост / ь»);
- overlap — текст двух рамок занимает одно место («универсальные» поверх следующего пункта);
- outside — текст выходит за фигуру, на которой стоит (карточка, панель, плашка);
- offslide — текст за границей слайда;
- ellipsis — текст кончается многоточием (обрезка запрещена, D-071).

Ошибка любого вида — колода не отдаётся, пока пайплайн её не исправит (core.pipeline).
Здесь же — доля зоны контента, занятой содержимым (fill_share): та же мера, что у галереи.
"""

import io
import math
import re
from dataclasses import dataclass

from pptx import Presentation

from core.fitting.metrics import line_height, overlong_words, text_width, wrap

EMU = 6350                         # EMU в единице сетки 1920×1080
ZONE = (80, 280, 1760, 640)        # зона контента [x, y, w, h] (05_LAYOUTS.md, 1)
FILL_MIN = 40.0                    # нижняя граница заполненности зоны контента, %
FILL_KINDS = {"bullets", "process", "conclusion", "metrics", "comparison"}
FILL_EXEMPT = {"metrics.big_number"}       # градиентная панель на всю зону — так задумано
DECOR_PREFIXES = ("decor", "header:")
SERVICE_NAMES = ("footer", "page_number", "gallery_label")
TOLERANCE = 1.0                    # единиц сетки: округление EMU
_ELLIPSIS_RE = re.compile(r"(…|\.\.\.)\s*$")
_ITEM_RE = re.compile(r"^(?P<name>[^#:]+)#(?P<item>\d+)$")


@dataclass
class Issue:
    slide: int                      # номер слайда с 1
    shape: str                      # имя фигуры: «title», «text#2» (пункт 2), «left_point3»
    code: str                       # overflow | word | overlap | outside | offslide | ellipsis
    detail: str = ""

    @property
    def slot(self) -> tuple[str, int | None]:
        """(имя элемента макета, номер пункта с 0 или None)."""
        m = _ITEM_RE.match(self.shape)
        if m:
            return m.group("name"), int(m.group("item")) - 1
        return self.shape.split(":")[0], None

    def __str__(self) -> str:
        return f"слайд {self.slide}: {self.shape} — {self.code}" + (f" ({self.detail})" if self.detail else "")


def units(v) -> float:
    return int(v) / EMU


def box_of(shape) -> tuple[float, float, float, float]:
    return units(shape.left), units(shape.top), units(shape.width), units(shape.height)


def is_decor(shape) -> bool:
    return shape.name.startswith(DECOR_PREFIXES)


def _runs(paragraph) -> list[tuple[str, int, bool]]:
    return [(r.text, round(r.font.size.pt) if r.font.size else 18, bool(r.font.bold)) for r in paragraph.runs]


def _paragraph_lines(paragraph, width: float, wrap_on: bool) -> list[tuple[float, float]]:
    """Строки абзаца: (ширина, высота). Абзац из нескольких фрагментов разного кегля
    (значение + единица) не переносится — это одна строка."""
    runs = _runs(paragraph)
    if not runs or not paragraph.text.strip():
        return []
    pt0 = runs[0][1]
    lh = paragraph.line_spacing.pt * 2 if getattr(paragraph.line_spacing, "pt", None) else line_height(pt0)
    if len(runs) > 1 or not wrap_on:
        return [(sum(text_width(t, pt, b) for t, pt, b in runs), lh)]
    text, pt, bold = runs[0]
    return [(text_width(line, pt, bold), lh) for line in wrap(text, width, pt, bold)]


def text_lines(shape) -> list[tuple[float, float, float, float]]:
    """Прямоугольники строк текста фигуры [x, y, w, h] по метрикам шрифта и якорю рамки."""
    x, y, w, h = box_of(shape)
    tf = shape.text_frame
    wrap_on = tf.word_wrap is not False
    lines = []
    for p in tf.paragraphs:
        for lw, lh in _paragraph_lines(p, w, wrap_on):
            lines.append((lw, lh, p.alignment))
    total = sum(lh for _, lh, _ in lines)
    anchor = tf.vertical_anchor
    top = y + {3: (h - total) / 2, 4: h - total}.get(int(anchor) if anchor is not None else 1, 0)
    out = []
    for lw, lh, align in lines:
        a = int(align) if align is not None else 1
        lx = x + (w - lw) / 2 if a == 2 else x + w - lw if a == 3 else x
        out.append((lx, top, lw, lh))
        top += lh
    return out


def _extent(rects):
    x0 = min(r[0] for r in rects)
    y0 = min(r[1] for r in rects)
    x1 = max(r[0] + r[2] for r in rects)
    y1 = max(r[1] + r[3] for r in rects)
    return x0, y0, x1, y1


def _intersects(a, b, pad: float = TOLERANCE) -> bool:
    return a[0] < b[2] - pad and b[0] < a[2] - pad and a[1] < b[3] - pad and b[1] < a[3] - pad


def _inside(inner, outer, pad: float = TOLERANCE) -> bool:
    return (inner[0] >= outer[0] - pad and inner[1] >= outer[1] - pad
            and inner[2] <= outer[2] + pad and inner[3] <= outer[3] + pad)


def _text_shapes(slide_obj):
    for shape in slide_obj.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        if len(shape.text_frame.text.strip()) == 1 and not shape.text_frame.text.strip().isalnum():
            continue                                            # знак-украшение: кавычка, галочка
        if shape.name in SERVICE_NAMES or is_decor(shape):
            continue
        yield shape


def _containers(slide_obj):
    """Фигуры-подложки, на которых стоит текст: карточки, панели, плашки, полосы (не декор)."""
    out = []
    for shape in slide_obj.shapes:
        if is_decor(shape) or shape.has_text_frame and shape.text_frame.text.strip():
            continue
        if shape.shape_type in (13, 3) or getattr(shape, "has_chart", False) and shape.has_chart:
            continue                                            # картинки и диаграммы
        x, y, w, h = box_of(shape)
        if w < 40 or h < 40:
            continue                                            # линии, полоски, точки
        out.append((shape.name, (x, y, x + w, y + h)))
    return out


def slide_issues(slide_obj, index: int) -> list[Issue]:
    issues: list[Issue] = []
    texts = []
    for shape in _text_shapes(slide_obj):
        x, y, w, h = box_of(shape)
        tf = shape.text_frame
        lines = text_lines(shape)
        if not lines:
            continue
        used = sum(r[3] for r in lines)
        if used > h + TOLERANCE:
            issues.append(Issue(index, shape.name, "overflow", f"{used:.0f} > {h:.0f}"))
        wrap_on = tf.word_wrap is not False
        for p in tf.paragraphs:
            runs = _runs(p)
            if len(runs) == 1 and wrap_on:
                long = overlong_words(runs[0][0], w, runs[0][1], runs[0][2])
                if long:
                    issues.append(Issue(index, shape.name, "word", long[0][:40]))
            elif any(lw > w * 1.0 + TOLERANCE for lw, _ in _paragraph_lines(p, w, wrap_on)):
                issues.append(Issue(index, shape.name, "overflow", "строка шире рамки"))
        if _ELLIPSIS_RE.search(tf.text.strip()):
            issues.append(Issue(index, shape.name, "ellipsis"))
        ext = _extent(lines)
        if ext[0] < -TOLERANCE or ext[1] < -TOLERANCE or ext[2] > 1920 + TOLERANCE or ext[3] > 1080 + TOLERANCE:
            issues.append(Issue(index, shape.name, "offslide"))
        texts.append((shape.name, (x, y, x + w, y + h), ext))

    # Наложение: текст двух рамок в одном месте
    for i in range(len(texts)):
        for j in range(i + 1, len(texts)):
            if _intersects(texts[i][2], texts[j][2]):
                later = texts[j][0]
                issues.append(Issue(index, later, "overlap", f"с {texts[i][0]}"))

    # Текст за фигурой, на которой стоит: рамка текста лежит в подложке, а текст — выходит за неё
    containers = _containers(slide_obj)
    for name, frame, ext in texts:
        hosts = [c for c in containers if _inside(frame, c[1])]
        if not hosts:
            continue
        host = min(hosts, key=lambda c: (c[1][2] - c[1][0]) * (c[1][3] - c[1][1]))
        if not _inside(ext, host[1]):
            issues.append(Issue(index, name, "outside", host[0]))
    return issues


def check_pptx(pptx: bytes) -> list[Issue]:
    """Все ошибки вёрстки готового файла."""
    prs = Presentation(io.BytesIO(pptx))
    issues: list[Issue] = []
    for i, slide_obj in enumerate(prs.slides, start=1):
        issues += slide_issues(slide_obj, i)
    return issues


# ── Заполненность зоны контента ─────────────────────────────────────────────

CELL = 10


def fill_share(slide_obj, zone=ZONE) -> float:
    """Доля зоны контента (сетка 10×10), занятой фигурами, картинками, диаграммами и текстом."""
    zx, zy, zw, zh = zone
    cols, rows = int(zw // CELL), int(zh // CELL)
    covered = set()

    def mark(x, y, w, h):
        c0, c1 = max(0, int((x - zx) // CELL)), min(cols, int(math.ceil((x + w - zx) / CELL)))
        r0, r1 = max(0, int((y - zy) // CELL)), min(rows, int(math.ceil((y + h - zy) / CELL)))
        for r in range(r0, r1):
            for c in range(c0, c1):
                covered.add((r, c))

    for shape in slide_obj.shapes:
        if is_decor(shape) or shape.name in ("title", "footnote", *SERVICE_NAMES):
            continue
        if shape.has_text_frame and not shape.text_frame.text.strip() and shape.shape_type == 17:
            continue
        if shape.shape_type == 17 or (shape.has_text_frame and shape.name.endswith((":number",))):
            for ext in text_lines(shape):
                mark(*ext)
            continue
        mark(*box_of(shape))
    return 100.0 * len(covered) / (cols * rows)
