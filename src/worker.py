"""
ARQ worker — фоновый пайплайн генерации презентации.

Раньше вся генерация (LLM → картинки → HTML → PDF) шла синхронно внутри
aiogram-хендлера в main.py и блокировала бота на 60-90 секунд: пока один
пользователь ждёт свою презентацию, бот не может ответить на /start или
/plan никому другому. Теперь main.py только кладёт job в очередь Redis
(arq_pool.enqueue_job) и сразу отвечает — вся тяжёлая работа происходит
здесь, в отдельном процессе.

Запуск: arq worker.WorkerSettings

Задачи (docs/design/08_MIGRATION.md, 2):
- generate_deck_job(deck_id) — новый движок (src/core): PPTX + PDF, доклад по теме
  и по материалу; колоды бота (файлы — в Telegram) и REST API (файлы — в хранилище
  core.storage.files, статус и скачивание — через API); питч-дек идёт в старый
  движок до сессии 5 (D-038, D-054);
- deliver_webhook_job(deck_id) — webhook колоды API (api.webhook);
- delete_message_job(chat_id, message_id) — удаление сообщения с ключом API (bot/admin.py);
- notify_admins_job(text) — сообщение владельцам: клиент API израсходовал 80% / 100% лимита (api.limit_alerts);
- generate_presentation_job — старый движок (HTML → PDF), удаляется в сессии 6.

Пайплайн старой задачи:
  extract_content (если source_type != TOPIC)
    → generate_presentation_structure
    → fetch_images_for_slides
    → render_presentation
    → html_to_pdf
    → загрузка в S3/MinIO
    → отправка пользователю в Telegram

Задача 3 (content_extractor) намеренно вызывается ЗДЕСЬ, а не в aiogram-
хендлере — парсинг pdf/docx/pptx или скачивание URL может занять секунды,
и это ровно та работа, которую мы не хотим держать в бот-процессе.
"""

import asyncio
import logging
import time

from aiogram import Bot
from aiogram.types import BufferedInputFile, InlineKeyboardButton, InlineKeyboardMarkup
from arq.connections import RedisSettings

from config import settings
from schemas.presentation import UserRequest, PresentationType, SlideLayout, SourceMode, _slide_has_required_content
from generation.content_extractor import extract_from_document, extract_from_url
from generation.llm import generate_presentation_structure, _default_slide_count
from generation.llm_models import track_usage
from generation.image_fetcher import fetch_images_for_slides
from generation.template_engine import render_presentation
from generation.pdf_renderer import html_to_pdf, get_renderer, shutdown_renderer
from generation.storage import upload_pdf
from generation.uploads import UploadNotFound, close_uploads, delete_upload, get_upload
from db.session import init_db, close_db, get_session, get_or_create_user, record_presentation, count_generation
from logging_setup import deck_id_var, setup_logging, stage
from bot import delivery
from core.models.request import DeckRequest
from core.models.theme import load_theme
from core.pipeline import DeckError, generate_deck
from core.storage import decks as deck_store
from core.storage import files as deck_files
from core.storage.progress import set_progress
from core.storage.redis import close_redis
from api.webhook import deliver_webhook_job
from bot.admin import delete_message_job, notify_admins_job

setup_logging()
logger = logging.getLogger(__name__)

JOB_TIMEOUT_SECONDS = 180  # LLM + картинки + PDF с запасом; см. WorkerSettings.job_timeout
# Новый движок: дедлайн колоды 110 с (DECK_DEADLINE_SECONDS), job_timeout — страховка
# от зависания кода (03_ARCHITECTURE.md, 4.2); ARQ берёт общий job_timeout воркера.
MAX_CONCURRENT_JOBS = 5    # нагрузочный критерий Sprint 1: 5 параллельных не блокируют бота

# Layout'ы, которые реально умеет рисовать templates/doklad/template.html
# (см. LAYOUT_CATALOG в llm.py). SlideLayout — общий enum на все типы
# презентаций, так что модель технически может вернуть "problem"/"market"/
# "team"/"competition"/"why_now" — их шаблон не рисует. Сам шаблон на такой
# случай не остаётся пустым (см. {% else %} в template.html), но нам нужно
# видеть в логах, если модель всё же выходит за пределы каталога.
DOKLAD_SUPPORTED_LAYOUTS = frozenset({
    SlideLayout.TITLE, SlideLayout.BULLETS, SlideLayout.METRICS,
    SlideLayout.TWO_COLUMN, SlideLayout.DIAGRAM, SlideLayout.QUOTE,
    SlideLayout.TIMELINE, SlideLayout.IMAGE_FULL, SlideLayout.IMAGE_HERO,
    SlideLayout.CLOSING,
})


def _kb_after_pdf() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Новая презентация", callback_data="action:new")],
    ])


async def generate_presentation_job(
    ctx: dict,
    *,
    job_id: str,
    chat_id: int,
    user_id: int,
    status_message_id: int | None,
    request_data: dict,
    document_ref: str | None = None,
    document_mime_type: str | None = None,
    # Старый контракт: байты прямо в аргументах. Оставлен, чтобы задачи,
    # поставленные в очередь до деплоя, не упали. Удалить в следующем релизе.
    document_bytes: bytes | None = None,
    urls: list[str] | None = None,
    watermark: bool = True,
    color_scheme: str = "light",
) -> dict:
    bot: Bot = ctx["bot"]
    # deck_id попадает во все строки лога этой задачи (logging_setup.JsonFormatter)
    deck_id_var.set(job_id)
    timings: dict[str, int] = {}
    job_started = time.perf_counter()

    async def _status(text: str) -> None:
        if status_message_id is None:
            return
        try:
            await bot.edit_message_text(text, chat_id=chat_id, message_id=status_message_id)
        except Exception:
            # Сообщение могло быть удалено пользователем — это не повод валить job
            logger.debug("Could not edit status message", extra={"job_id": job_id})

    try:
        # UserRequest.raw_text is required by a model validator whenever
        # source_type != TOPIC — so we can't construct the model until
        # extraction has filled it in. Work on the raw dict first, extract,
        # then validate once into UserRequest.
        request_data = dict(request_data)
        source_type = request_data.get("source_type", "topic")

        if source_type != "topic" and not request_data.get("raw_text"):
            await _status("⚙️ Читаю материал...")
            try:
                with stage("ingest", timings):
                    if document_ref is not None:
                        document_bytes = await get_upload(document_ref)
                    raw_text = await _extract_raw_text(job_id, document_bytes, document_mime_type, urls)
            except UploadNotFound:
                # Задача пролежала в очереди дольше TTL файла — просим прислать заново
                logger.warning("Uploaded document expired or missing", extra={"ref": document_ref})
                await _status("❌ Файл устарел — пришлите его ещё раз через /new")
                return {"status": "failed", "job_id": job_id, "reason": "upload_not_found"}
            request_data["raw_text"] = raw_text

        request = UserRequest.model_validate(request_data)

        await _status("⚙️ Пишу текст слайдов...")
        with stage("llm", timings), track_usage() as usage:
            presentation = await generate_presentation_structure(request)
        # Токены и стоимость по колоде (ТЗ 4.7); рассуждения входят в output_tokens
        logger.info("LLM usage", extra={"model": settings.openai_model, **usage.as_dict()})

        if request.presentation_type == PresentationType.DOKLAD:
            # Диагностика написана прямо в тексте сообщения — это осталось с
            # тех пор, когда логи теряли extra={...}. Теперь логи JSON и extra
            # печатаются (logging_setup), но текст оставлен как есть.
            expected_count = request.slide_count_hint or _default_slide_count(request.presentation_type)
            logger.info("Doklad generated", extra={
                "source_genre": presentation.meta.source_genre,
                "source_mode": request.source_mode.value if request.source_mode else None,
                "slide_count": presentation.slide_count, "slide_count_hint": expected_count,
            })
            # В strict меньше слайдов — законный исход (материала мало), не обрезка
            if (presentation.slide_count != expected_count
                    and not _material_shortage_note(request, presentation.slide_count)):
                logger.warning(
                    f"DOKLAD job={job_id}: LLM returned {presentation.slide_count} slides, "
                    f"expected {expected_count} — likely response truncation (see max_tokens "
                    f"in generate_presentation_structure), slides past the cut come out with "
                    f"missing fields"
                )

            def _slide_debug(s) -> str:
                return (
                    f"#{s.index} layout={s.layout.value} "
                    f"title={s.title!r} subtitle={s.subtitle!r} "
                    f"body_text_len={len(s.body_text) if s.body_text else 0} "
                    f"bullets={len(s.bullets)} metrics={len(s.metrics)} "
                    f"two_column={'yes' if s.two_column else 'no'} "
                    f"timeline_items={len(s.timeline_items)} "
                    f"mermaid_code={'yes' if s.mermaid_code else 'no'} "
                    f"image_query={s.image_query!r}"
                )

            unsupported = [s for s in presentation.slides if s.layout not in DOKLAD_SUPPORTED_LAYOUTS]
            if unsupported:
                logger.warning(
                    f"DOKLAD job={job_id}: LLM returned layout(s) outside the doklad "
                    f"catalog — rendered via the generic fallback block instead of a "
                    f"dedicated layout. " + "; ".join(_slide_debug(s) for s in unsupported)
                )

            # Слайд может иметь ПОДДЕРЖИВАЕМЫЙ layout (прошёл проверку выше) и
            # всё равно оказаться визуально пустым или полупустым — либо ни
            # одного поля вообще, либо title+subtitle есть, а само "тело"
            # (bullets/metrics/two_column/mermaid_code/...), которое рисует
            # именно ЭТОТ layout, пустое (реальный прод-кейс: title="Логика
            # управления требованиями", subtitle="Понимание процесса",
            # layout=two_column, two_column=None). Используем ТОТ ЖЕ чек,
            # что раньше жёстко ронял валидацию (см. schemas/presentation.py)
            # — теперь только логируем, не блокируя выдачу презентации:
            # блокировка на практике била по надёжности сильнее, чем сам
            # баг (модель может стабильно повторять пустой слайд все 3
            # попытки retry, и пользователь не получал вообще ничего).
            content_gap_slides = [
                s for s in presentation.slides
                if s.layout not in (SlideLayout.TITLE, SlideLayout.CLOSING)
                and not _slide_has_required_content(s)
            ]
            if content_gap_slides:
                logger.warning(
                    f"DOKLAD job={job_id}: slide(s) have a supported layout but are "
                    f"missing the specific field that layout renders — will show as a "
                    f"blank or near-blank page. " + "; ".join(_slide_debug(s) for s in content_gap_slides)
                )

        # ФИО/группа докладчика — не от модели, а из профиля пользователя в
        # боте (см. main.py, "👤 Профиль"). Только для DOKLAD: это единственный
        # шаблон, который их рендерит (title/closing слайды).
        if request.presentation_type == PresentationType.DOKLAD:
            async with get_session() as session:
                if session:
                    profile_user = await get_or_create_user(session, user_id)
                    if profile_user.author_name or profile_user.author_group:
                        presentation = presentation.model_copy(update={
                            "meta": presentation.meta.model_copy(update={
                                "author_name": profile_user.author_name,
                                "author_group": profile_user.author_group,
                            })
                        })

        await _status("⚙️ Подбираю изображения...")
        with stage("images", timings):
            image_urls = await fetch_images_for_slides(presentation.slides)

        await _status("⚙️ Собираю дизайн...")
        with stage("render_html", timings):
            html = render_presentation(
                presentation,
                image_urls=image_urls,
                watermark=watermark,
                color_scheme=color_scheme,
            )

        await _status("⚙️ Рендерю PDF...")
        has_mermaid = any(s.layout.value == "diagram" for s in presentation.slides)
        with stage("render_pdf", timings):
            pdf_bytes = await html_to_pdf(html, has_mermaid=has_mermaid, expected_pages=presentation.slide_count)

        with stage("upload", timings):
            await upload_pdf(job_id, pdf_bytes)

        # JSON презентации сохраняем вместе с записью (ТЗ, раздел 8, п. 14):
        # разбор жалоб, аналитика, будущий переэкспорт без новой генерации.
        # Длительности — всё, что посчитано к этому моменту (без deliver).
        with stage("db_record", timings):
            async with get_session() as session:
                if session:
                    user = await get_or_create_user(session, user_id)
                    await record_presentation(
                        session,
                        user=user,
                        topic=request.topic,
                        presentation_type=request.presentation_type.value,
                        audience=request.audience.value,
                        language=request.language,
                        slide_count=presentation.slide_count,
                        has_brief=bool(request.extra_instructions),
                        watermark=watermark,
                        job_id=job_id,
                        spec=presentation.model_dump(mode="json"),
                        durations_ms=dict(timings),
                    )

        if status_message_id is not None:
            try:
                await bot.delete_message(chat_id=chat_id, message_id=status_message_id)
            except Exception:
                pass

        safe_title = "".join(
            c if c.isalnum() or c in " _-" else "_"
            for c in presentation.meta.title
        )[:40]

        caption = (
            f"✨ <b>{presentation.meta.title}</b>\n"
            f"{presentation.slide_count} слайдов"
        )
        short_note = _material_shortage_note(request, presentation.slide_count)
        if short_note:
            caption += f"\n\n{short_note}"
        if watermark:
            caption += "\n\n<i>Бесплатная версия · Уберите водяной знак в /plan</i>"

        with stage("deliver", timings):
            await bot.send_document(
                chat_id=chat_id,
                document=BufferedInputFile(pdf_bytes, filename=f"{safe_title}.pdf"),
                caption=caption,
                parse_mode="HTML",
                reply_markup=_kb_after_pdf(),
            )

        logger.info(
            "Job completed",
            extra={
                "job_id": job_id, "chat_id": chat_id, "slide_count": presentation.slide_count,
                "total_ms": int((time.perf_counter() - job_started) * 1000), "durations_ms": timings,
            },
        )
        return {"status": "ok", "job_id": job_id, "slide_count": presentation.slide_count}

    except Exception as e:
        logger.exception(
            "Job failed",
            extra={
                "job_id": job_id, "chat_id": chat_id, "error": str(e),
                "total_ms": int((time.perf_counter() - job_started) * 1000), "durations_ms": timings,
            },
        )
        await _status("❌ Что-то пошло не так. Попробуйте ещё раз через /new")
        raise

    finally:
        # Файлы пользователя не храним дольше обработки (ТЗ 6.5)
        if document_ref is not None:
            await delete_upload(document_ref)


async def generate_deck_job(ctx: dict, deck_id: str, request: dict | None = None) -> dict:
    """Новый движок: колода по deck_id (D-034). request передаётся, только если у бота
    не было БД и строку decks создать не удалось."""
    bot: Bot = ctx["bot"]
    deck_id_var.set(deck_id)
    raw = request or await deck_store.load_request(deck_id)
    if raw is None:
        logger.error("Deck request not found", extra={"deck_id": deck_id})
        return {"status": "failed", "deck_id": deck_id, "reason": "not_found"}
    req = DeckRequest.model_validate(raw)
    chat_id, user_id = req.client.chat_id, req.client.user_id
    status = delivery.StatusMessage(bot, chat_id, req.client.status_message_id, deck_id)
    stage_seen: set[str] = set()

    is_api = req.client.kind == "api"

    async def progress(stage_name: str, done: int | None = None, total: int | None = None) -> None:
        if stage_name not in stage_seen:
            stage_seen.add(stage_name)
            await deck_store.set_stage(deck_id, stage_name)
        if is_api and total:
            await set_progress(deck_id, done or 0, total)
        await status(stage_name, done, total)

    async def webhook() -> None:
        if is_api and req.webhook_url and ctx.get("redis") is not None:
            try:
                await ctx["redis"].enqueue_job("deliver_webhook_job", deck_id)
            except Exception as e:
                logger.warning("Webhook not enqueued", extra={"error": str(e)})

    await deck_store.mark_processing(deck_id)
    started = time.perf_counter()
    material_ref = req.input.material.ref if req.input.material else None
    try:
        material = None
        if material_ref:
            try:
                material = await get_upload(material_ref)
            except UploadNotFound:
                raise DeckError("UPLOAD_EXPIRED", "uploaded material expired")
        result = await generate_deck(deck_id, req, progress, material=material)
    except DeckError as e:
        logger.error("Deck failed", extra={"error_code": e.code, "error": str(e)})
        await deck_store.fail(deck_id, e.code)
        await status.show(delivery.error_text(e.code, deck_id))
        await webhook()
        return {"status": "failed", "deck_id": deck_id, "error_code": e.code}
    except Exception as e:
        logger.exception("Deck crashed", extra={"error": str(e)})
        await deck_store.fail(deck_id, "INTERNAL")
        await status.show(delivery.error_text(None, deck_id))
        await webhook()
        raise
    finally:
        # Файлы пользователя не храним дольше обработки (ТЗ 6.5)
        if material_ref:
            await delete_upload(material_ref)

    spec = result.spec
    warnings = list(result.warnings) + (["pdf_failed"] if result.pdf is None else [])
    files = None
    if is_api:
        # Файлы колоды API — в хранилище: клиент скачивает их через API (D-058)
        try:
            with stage("store_files", result.durations_ms):
                files = await deck_files.put_deck_files(deck_id, {"pptx": result.pptx, "pdf": result.pdf})
        except Exception as e:
            logger.exception("Deck files not stored", extra={"error": str(e)})
            await deck_store.fail(deck_id, "STORAGE_FAILED", durations_ms=result.durations_ms,
                                  usage=result.usage.as_dict(), cost_rub=result.usage.cost_rub,
                                  degradations=[d.model_dump() for d in spec.degradations])
            await webhook()
            return {"status": "failed", "deck_id": deck_id, "error_code": "STORAGE_FAILED"}
    await deck_store.finish(
        deck_id, status="done", spec=spec.model_dump(mode="json"), durations_ms=result.durations_ms,
        usage=result.usage.as_dict(), cost_rub=result.usage.cost_rub,
        degradations=[d.model_dump() for d in spec.degradations], files=files, warnings=warnings,
    )
    if is_api:
        # Колода API готова, как только файлы в хранилище; лимит API — по строкам decks (api.store)
        await deck_store.mark_counted(deck_id)
        await webhook()

    notes = delivery.material_notes(result.warnings) if req.mode == "material" else []
    text = delivery.caption(spec.meta.title, len(spec.slides), load_theme(spec.meta.theme_id).name.get("ru", ""),
                            req.watermark, result.pdf is None, notes)
    if chat_id is not None:
        with stage("deliver", result.durations_ms):
            await delivery.send_files(bot, chat_id, spec.meta.title, result.pptx, result.pdf, text)
        await status.delete()
        # Лимит списывается только после успешной отправки (02_CJM.md, 1.2, шаг 8)
        async with get_session() as session:
            if session and user_id is not None:
                await count_generation(session, user_id)
        await deck_store.mark_counted(deck_id)

    logger.info("Deck delivered", extra={
        "slides": len(spec.slides), "cost_rub": result.usage.cost_rub,
        "total_ms": int((time.perf_counter() - started) * 1000), "durations_ms": result.durations_ms,
    })
    return {"status": "ok", "deck_id": deck_id, "slides": len(spec.slides), "cost_rub": result.usage.cost_rub}


def _material_shortage_note(request: UserRequest, slide_count: int) -> str | None:
    """«Только мой материал»: слайдов не больше выбранного, и если модель
    сделала меньше — говорим пользователю, что материала хватило на X (ТЗ 3.3.3)."""
    if request.source_type.value == "topic" or not request.slide_count_hint:
        return None
    if (request.source_mode or SourceMode.STRICT) != SourceMode.STRICT:
        return None
    if slide_count >= request.slide_count_hint:
        return None
    return (
        f"В вашем материале хватило на {slide_count} слайдов из {request.slide_count_hint} — "
        f"добавьте текст, если нужно больше."
    )


async def _extract_raw_text(
    job_id: str,
    document_bytes: bytes | None,
    document_mime_type: str | None,
    urls: list[str] | None,
) -> str | None:
    if document_bytes is not None:
        return await extract_from_document(document_bytes, document_mime_type)
    if urls:
        raw_text, error_urls = await extract_from_url(urls)
        if error_urls:
            logger.warning(
                "Some URLs failed to extract",
                extra={"job_id": job_id, "error_urls": error_urls},
            )
        return raw_text
    return None


async def startup(ctx: dict) -> None:
    # CLI arq после импорта модуля ставит свой текстовый хендлер — переопределяем
    setup_logging()
    if not settings.openai_api_key or not settings.telegram_bot_token:
        raise RuntimeError("OPENAI_API_KEY и TELEGRAM_BOT_TOKEN обязательны для воркера")
    ctx["bot"] = Bot(token=settings.telegram_bot_token)
    await init_db()
    await get_renderer()
    # Первый запуск LibreOffice в контейнере медленный (холодный диск) — прогреваем
    # в фоне, чтобы первая колода не потратила на это бюджет CONVERT
    ctx["lo_warmup"] = asyncio.create_task(_warmup_libreoffice())
    logger.info("ARQ worker started", extra={"max_jobs": MAX_CONCURRENT_JOBS})


async def _warmup_libreoffice() -> None:
    import io
    from pptx import Presentation
    from core.render.pdf.convert import pptx_to_pdf

    buf = io.BytesIO()
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(buf)
    started = time.perf_counter()
    try:
        await pptx_to_pdf(buf.getvalue(), timeout=120)
        logger.info("LibreOffice warmed up", extra={"duration_ms": int((time.perf_counter() - started) * 1000)})
    except Exception as e:
        logger.error("LibreOffice warmup failed — PDF of the new engine will not work", extra={"error": str(e)})


async def shutdown(ctx: dict) -> None:
    await shutdown_renderer()
    await close_uploads()
    await close_redis()
    await close_db()
    bot: Bot | None = ctx.get("bot")
    if bot is not None:
        await bot.session.close()
    logger.info("ARQ worker stopped")


class WorkerSettings:
    # generate_presentation_job (старый движок) — для питч-дека до сессии 5
    # и для задач, поставленных до деплоя (08_MIGRATION.md, 2)
    functions = [generate_deck_job, deliver_webhook_job, delete_message_job, notify_admins_job,
                 generate_presentation_job]
    on_startup = startup
    on_shutdown = shutdown
    redis_settings = RedisSettings.from_dsn(settings.redis_url)
    job_timeout = JOB_TIMEOUT_SECONDS
    max_jobs = MAX_CONCURRENT_JOBS
    # job_id не задан пользователем -> arq сам сгенерирует; мы передаём свой
    # job_id как аргумент задачи отдельно (для статус-сообщений/S3-ключа),
    # см. main.py: enqueue_job(..., _job_id=job_id, job_id=job_id, ...)
