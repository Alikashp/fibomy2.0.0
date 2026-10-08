"""Стресс-галерея макетов (07_TESTING.md, 2; ТЗ 7.3): каждый вариант × каждая тема ×
3 заполнения (min, typical, max) → PPTX → проверки → PDF → PNG.

    python tests/gallery/build_gallery.py OUT_DIR [--themes a,b] [--fixtures typical] [--variants x,y]

В OUT_DIR: gallery.pdf (книга: тема → заполнение → вариант; подпись внизу каждой страницы),
sheet_<тема>.png (миниатюры заполнения typical), report.md и results.json (проверки).
Код выхода 1 — есть нарушения (job «gallery» в workflow «tests» красный).
Без LLM и сети; LibreOffice обязателен.
"""

import argparse
import asyncio
import io
import json
import logging
import statistics
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import pypdfium2 as pdfium  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Pt  # noqa: E402

import checks as C  # noqa: E402
import fixtures as F  # noqa: E402
from core.fitting.fitter import fit_deck  # noqa: E402
from core.models.deck import DeckSpec, Meta, Slide  # noqa: E402
from core.models.ids import slide_id  # noqa: E402
from core.models.layout import load_layouts  # noqa: E402
from core.models.theme import all_theme_ids, load_theme  # noqa: E402
from core.render.pdf.convert import pptx_to_pdf  # noqa: E402
from core.render.pptx.renderer import render_pptx  # noqa: E402

DECK_ID = "dk_01J00000000000000000000000"
CHUNK = 20
FIXTURE_ORDER = ("typical", "min", "max")


def build_slides(variants: list[str], fixture: str) -> tuple[list[Slide], dict[str, bytes]]:
    layouts = load_layouts()
    slides, images = [], {}
    for i, vid in enumerate(variants, start=1):
        layout = layouts[vid]
        title, content, n = F.content_for(layout, fixture, seed=i * 7)
        image = None
        if layout.image_slot:
            image = f"img_{i:08d}"
            images[image] = F.placeholder_image(i)
        slides.append(Slide(id=slide_id(i), index=i, kind=layout.kind, variant=vid, title=title, content=content,
                            data=F.chart_data(vid, fixture), image=image,
                            footnote="Оценочные данные — проверьте перед показом"
                            if fixture == "max" and layout.kind not in ("title", "closing") else None))
        if layout.kind == "chart_share":
            from core.planning.datasets import legend
            slides[-1].content["legend"] = legend(slides[-1].data)
    return slides, images


def deck(theme_id: str, slides: list[Slide], title: str) -> DeckSpec:
    while len(slides) < 4:
        slides = slides + [slides[-1].model_copy(update={"id": slide_id(len(slides) + 1), "index": len(slides) + 1})]
    for i, s in enumerate(slides, start=1):
        s.index, s.id = i, slide_id(i)
    return DeckSpec(id=DECK_ID, meta=Meta(title=title, language="ru", presentation_type="doklad", audience="general",
                                          mode="topic", theme_id=theme_id, seed=1), slides=slides)


def label(pptx: bytes, labels: list[str]) -> bytes:
    """Подпись «вариант · тема · заполнение» внизу каждой страницы галереи (только в галерее)."""
    prs = Presentation(io.BytesIO(pptx))
    for slide_obj, text in zip(prs.slides, labels):
        box = slide_obj.shapes.add_textbox(0, prs.slide_height - 230000, prs.slide_width, 200000)
        box.name = "gallery_label"
        p = box.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = text
        run.font.size = Pt(9)
        run.font.name = "Arial"
        from pptx.dml.color import RGBColor
        run.font.color.rgb = RGBColor(0x9A, 0xA0, 0xA8)
        from pptx.enum.text import PP_ALIGN
        p.alignment = PP_ALIGN.CENTER
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


async def run(out: Path, themes: list[str], fixtures: list[str], variants: list[str]) -> int:
    out.mkdir(parents=True, exist_ok=True)
    book = pdfium.PdfDocument.new()
    results, failures = [], 0
    sheets: dict[str, list[Image.Image]] = {}
    layouts = load_layouts()
    for theme_id in themes:
        theme = load_theme(theme_id)
        for fixture in fixtures:
            slides, images = build_slides(variants, fixture)
            fit_errors_by_variant = {}
            # Тема титула — meta.title колоды: каждый вариант титула — своей колодой со своей темой
            titles = [s for s in slides if s.kind == "title"]
            rest = [s for s in slides if s.kind != "title"]
            chunks = [[t] for t in titles] + [rest[i:i + CHUNK] for i in range(0, len(rest), CHUNK)]
            for start, part in enumerate(chunks):
                chunk = [s.model_copy(deep=True) for s in part]
                spec = deck(theme_id, chunk, chunk[0].title if chunk[0].kind == "title" else "Галерея")
                fit_deck(spec, theme)
                for d in spec.degradations:
                    if d.reason.startswith(("truncated", "overflow")):
                        sl = next(s for s in spec.slides if s.id == d.slide_id)
                        fit_errors_by_variant.setdefault(sl.variant, []).append(d.reason)
                labels = [f"{s.variant} · {theme_id} · {fixture}" for s in spec.slides]
                pptx = label(render_pptx(spec, images=images), labels)
                pdf_bytes = await pptx_to_pdf(pptx, timeout=180)
                (out / "pptx").mkdir(exist_ok=True)
                (out / "pptx" / f"{theme_id}_{fixture}_{start + 1}.pptx").write_bytes(pptx)
                pdf = pdfium.PdfDocument(pdf_bytes)
                prs = Presentation(io.BytesIO(pptx))
                for page_no, (slide, slide_obj) in enumerate(zip(spec.slides, prs.slides)):
                    if page_no >= len(chunk):
                        break
                    png = pdf[page_no].render(scale=640 / pdf[page_no].get_width()).to_pil()
                    layout = layouts[slide.variant]
                    pictures = [C.box_of(sh) for sh in slide_obj.shapes if sh.name.startswith("image:")]
                    errors = C.geometry_errors(slide_obj, theme) + C.chart_errors(slide_obj)
                    if fixture in ("min", "typical"):
                        errors += [f"fit: {r}" for r in fit_errors_by_variant.get(slide.variant, [])]
                    elif any(r.startswith("overflow") for r in fit_errors_by_variant.get(slide.variant, [])):
                        errors += [f"fit: {r}" for r in fit_errors_by_variant[slide.variant] if r.startswith("overflow")]
                    share = C.color_share(png, theme_id, pictures)
                    threshold = C.color_threshold(slide.kind, slide.variant)
                    if share < threshold and fixture != "min":   # min — короткие тексты: только в отчёте
                        errors.append(f"доля цвета {share:.1f}% < {threshold:.0f}%")
                    fill = C.fill_share(slide_obj) if slide.kind in C.FILL_KINDS and slide.variant not in C.FILL_EXEMPT else None
                    if fill is not None and fixture == "typical" and not C.FILL_RANGE[0] <= fill <= C.FILL_RANGE[1]:
                        errors.append(f"заполненность {fill:.0f}% вне {C.FILL_RANGE[0]:.0f}–{C.FILL_RANGE[1]:.0f}%")
                    results.append({"theme": theme_id, "fixture": fixture, "variant": slide.variant,
                                    "color_share": round(share, 1), "fill": None if fill is None else round(fill, 1),
                                    "errors": errors})
                    failures += bool(errors)
                    if fixture == "typical":
                        sheets.setdefault(theme_id, []).append(png)
                book.import_pages(pdf, list(range(len(chunk))))
    # Медиана доли цвета по «колоде» темы (заполнение typical) — ≥ 8%
    for theme_id in themes:
        shares = [r["color_share"] for r in results if r["theme"] == theme_id and r["fixture"] == "typical"]
        if shares and statistics.median(shares) < C.DECK_MEDIAN:
            results.append({"theme": theme_id, "fixture": "typical", "variant": "(медиана колоды)",
                            "color_share": statistics.median(shares), "fill": None,
                            "errors": [f"медиана доли цвета {statistics.median(shares):.1f}% < {C.DECK_MEDIAN:.0f}%"]})
            failures += 1
    book.save(str(out / "gallery.pdf"))
    for theme_id, pngs in sheets.items():
        contact_sheet(pngs, out / f"sheet_{theme_id}.png", theme_id)
    (out / "results.json").write_text(json.dumps(results, ensure_ascii=False, indent=1), encoding="utf-8")
    (out / "report.md").write_text(report(results, themes, fixtures, len(variants)), encoding="utf-8")
    print(report(results, themes, fixtures, len(variants)))
    return 1 if failures else 0


def contact_sheet(pngs: list[Image.Image], path: Path, title: str, cols: int = 5) -> None:
    w, h = 384, 216
    rows = (len(pngs) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * (w + 8) + 8, rows * (h + 8) + 40), "white")
    ImageDraw.Draw(sheet).text((8, 10), title, fill=(60, 60, 60),
                               font=ImageFont.truetype(str(ROOT / "fonts" / "LiberationSans-Bold.ttf"), 20))
    for i, im in enumerate(pngs):
        sheet.paste(im.resize((w, h)), (8 + (i % cols) * (w + 8), 40 + (i // cols) * (h + 8)))
    sheet.save(path, optimize=True)


def report(results: list[dict], themes: list[str], fixtures: list[str], n_variants: int) -> str:
    bad = [r for r in results if r["errors"]]
    lines = [f"## Стресс-галерея: {n_variants} вариантов × {len(themes)} темы × {len(fixtures)} заполнения "
             f"= {sum(1 for r in results if not r['variant'].startswith('('))} слайдов",
             "", f"Нарушений: **{len(bad)}**", ""]
    lines += ["| Тема | Доля цвета: медиана / минимум, % | Заполненность typical: мин–макс, % |", "|---|---|---|"]
    for t in themes:
        rs = [r for r in results if r["theme"] == t and not r["variant"].startswith("(")]
        shares = [r["color_share"] for r in rs if r["fixture"] == "typical"]
        fills = [r["fill"] for r in rs if r["fixture"] == "typical" and r["fill"] is not None]
        lines.append(f"| {t} | {statistics.median(shares):.1f} / {min(shares):.1f} | "
                     f"{min(fills):.0f}–{max(fills):.0f} |" if shares and fills else f"| {t} | — | — |")
    if bad:
        lines += ["", "| Вариант | Тема | Заполнение | Нарушения |", "|---|---|---|---|"]
        for r in bad:
            lines.append(f"| {r['variant']} | {r['theme']} | {r['fixture']} | {'; '.join(r['errors'])[:300]} |")
    return "\n".join(lines) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--themes", default=",".join(all_theme_ids()))
    ap.add_argument("--fixtures", default=",".join(FIXTURE_ORDER))
    ap.add_argument("--variants", default="")
    args = ap.parse_args()
    catalog = list(load_layouts())
    order = sorted(catalog, key=lambda v: (v.split(".")[0] != "title", v.split(".")[0] == "closing"))
    variants = [v for v in order if not args.variants or v in args.variants.split(",")]
    logging.disable(logging.WARNING)       # предупреждения fitter'а — в отчёте, а не в логе
    started = time.monotonic()
    code = asyncio.run(run(Path(args.out), args.themes.split(","), args.fixtures.split(","), variants))
    print(f"Готово за {time.monotonic() - started:.0f} с")
    return code


if __name__ == "__main__":
    sys.exit(main())
