"""FIT v0 (метрики, шаги кегля, обрезка) и RENDER/CONVERT: PPTX читается обратно, PDF собирается."""
import asyncio
import io
import re
import unittest

import _helpers  # noqa: F401 — путь к src/ и фиктивные ключи

from pptx import Presentation
from pptx.oxml.ns import qn

from core.fitting.fitter import fit_deck, fit_text, measure, truncate
from core.models.deck import DeckSpec, Meta, Slide
from core.models.ids import new_deck_id
from core.models.layout import capacity, load_layouts
from core.models.theme import load_theme
from core.render.pdf.convert import pptx_to_pdf, soffice_bin
from core.render.pptx.renderer import render_pptx

THEME = load_theme("graphite_light")
L = load_layouts()
RU = ("Растения превращают энергию солнечного света в энергию химических связей, и ею питаются все "
      "остальные организмы на планете. ") * 6


def make_spec(topic="Как работает фотосинтез", language="ru", watermark=False, bullets_text=None) -> DeckSpec:
    meta = Meta(title=topic, subtitle="Свет, вода и углекислый газ", language=language, presentation_type="doklad",
                audience="students", mode="topic", theme_id="graphite_light", seed=1,
                author={"name": "Иван Петров", "group": "Б-21"}, watermark=watermark)
    item = bullets_text or "Энергия света запускает реакции в хлоропластах."
    slides = [
        Slide(id="s01", index=1, kind="title", variant="title.cover_center", title=topic,
              content={"subtitle": meta.subtitle, "author": "Иван Петров · Б-21"}),
        Slide(id="s02", index=2, kind="statement", variant="statement.big_quote",
              title="Почти весь кислород в воздухе выделили растения", content={"body": "Пояснение."}),
        Slide(id="s03", index=3, kind="bullets", variant="bullets.cards_grid", title="Четыре условия фотосинтеза",
              content={"items": [{"heading": f"Условие {i}", "text": item, "icon": "bulb"} for i in range(4)]},
              footnote="Оценочные данные — проверьте перед показом"),
        Slide(id="s04", index=4, kind="conclusion", variant="conclusion.numbered_takeaways", title="Выводы-утверждения",
              content={"items": [{"text": "Вывод один."}, {"text": "Вывод два."}]}),
        Slide(id="s05", index=5, kind="closing", variant="closing.thanks_center", title="Спасибо за внимание",
              content={"title": "Спасибо за внимание", "body": None, "author": "Иван Петров · Б-21"}),
    ]
    return DeckSpec(id=new_deck_id(), meta=meta, slides=slides)


class Fitting(unittest.TestCase):

    def test_text_at_base_capacity_fits_base_style(self):
        """Формула вместимости — оценка сверху: русский текст длиной в вместимость помещается
        базовым стилем или на шаг ниже (±10% к таблицам 05_LAYOUTS.md)."""
        cases = [("bullets.cards_grid", 3, "item.text"), ("bullets.cards_grid", 4, "item.text"),
                 ("statement.big_quote", None, "statement"), ("conclusion.numbered_takeaways", 3, "item.text"),
                 ("title.cover_center", None, "subtitle")]
        for variant, n, slot in cases:
            e = L[variant].text_slots(n)[slot]
            text = RU[:capacity(e)].rsplit(" ", 1)[0]
            self.assertTrue(measure(e, text, e.style, THEME).fits, (variant, n, slot, len(text)))
            longer = RU[:int(capacity(e) * 1.4)]
            self.assertFalse(measure(e, longer, e.style, THEME).fits, (variant, n, slot))

    def test_style_steps_down_then_truncates(self):
        e = L["bullets.cards_grid"].text_slots(4)["item.text"]
        mid = RU[:170]
        self.assertEqual(fit_text(e, mid, THEME).style, "small")
        cut = truncate(e, RU, "small", THEME)
        self.assertTrue(cut.endswith("."))
        self.assertTrue(measure(e, cut, "small", THEME).fits)
        self.assertLess(len(cut), len(RU))

    def test_long_words_kk_and_de_wrap_by_chars(self):
        e = L["bullets.cards_grid"].text_slots(6)["item.heading"]
        for word in ("Qualitätssicherungsmaßnahmen", "Жауапкершіліктерініздің"):
            result = fit_text(e, word, THEME)
            self.assertTrue(result.fits, word)
            self.assertGreaterEqual(len(result.lines), 1)

    def test_deck_fit_keeps_topic_verbatim(self):
        topic = ("Управление требованиями заинтересованных сторон в ИТ-стартапе: от интервью до приоритизации "
                 "бэклога и проверки гипотез на пользователях, " * 2)[:200]
        spec = make_spec(topic=topic, bullets_text=RU)
        fit_deck(spec, THEME)
        self.assertEqual(spec.slides[0].title, topic)
        self.assertEqual(spec.meta.title, topic)
        self.assertIn(spec.slides[0].fit.styles["title"], ("h1", "h2"))
        bullets = spec.slides[2]
        self.assertTrue(bullets.fit.truncated)
        self.assertEqual(bullets.fit.styles["item.text"], "small")
        self.assertTrue(any(d.reason.startswith("truncated:") for d in spec.degradations))


class Render(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.spec = make_spec()
        fit_deck(cls.spec, THEME)
        cls.pptx = render_pptx(cls.spec)
        cls.prs = Presentation(io.BytesIO(cls.pptx))

    def texts(self, slide):
        return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]

    def test_slides_and_title_verbatim(self):
        self.assertEqual(len(self.prs.slides), 5)
        self.assertIn("Как работает фотосинтез", self.texts(self.prs.slides[0]))
        self.assertEqual(self.prs.slide_width, 12_192_000)

    def test_no_shadows_no_autofit_arial_lang(self):
        for slide in self.prs.slides:
            for sh in slide.shapes:
                self.assertIsNone(sh._element.find(qn("p:style")), sh.name)
                if sh.has_text_frame:
                    body = sh.text_frame._txBody.find(qn("a:bodyPr"))
                    self.assertIsNone(body.find(qn("a:normAutofit")))
                    self.assertIsNone(body.find(qn("a:spAutoFit")))
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            self.assertEqual(r.font.name, "Arial")
                            self.assertEqual(r._r.find(qn("a:rPr")).get("lang"), "ru-RU")

    def test_inside_slide(self):
        for slide in self.prs.slides:
            for sh in slide.shapes:
                self.assertTrue(sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= self.prs.slide_width
                                and sh.top + sh.height <= self.prs.slide_height, sh.name)

    def test_cards_icons_and_footer(self):
        slide = self.prs.slides[2]
        pictures = [sh for sh in slide.shapes if sh.shape_type == 13]
        self.assertEqual(len(pictures), 4)
        texts = self.texts(slide)
        self.assertIn("3 / 5", texts)
        self.assertIn("Оценочные данные — проверьте перед показом", texts)
        self.assertNotIn("Fibonacci.", texts)            # в PPTX водяного знака нет
        self.assertNotIn("1 / 5", self.texts(self.prs.slides[0]))  # на титуле номера нет

    def test_watermark_copy(self):
        prs = Presentation(io.BytesIO(render_pptx(self.spec, watermark=True)))
        for slide in prs.slides:
            self.assertIn("Fibonacci.", self.texts(slide))

    def test_metadata(self):
        cp = self.prs.core_properties
        self.assertEqual((cp.title, cp.author, cp.last_modified_by), ("Как работает фотосинтез", "Иван Петров", "Fibonacci AI"))
        app = next(p for p in self.prs.part.package.iter_parts() if str(p.partname) == "/docProps/app.xml")
        self.assertIn(b"<Application>Fibonacci AI</Application>", app.blob)

    def test_kazakh_lang_tag(self):
        spec = make_spec(language="kk")
        prs = Presentation(io.BytesIO(render_pptx(spec)))
        run = next(r for sh in prs.slides[1].shapes if sh.has_text_frame
                   for p in sh.text_frame.paragraphs for r in p.runs)
        self.assertEqual(run._r.find(qn("a:rPr")).get("lang"), "kk-KZ")


@unittest.skipUnless(soffice_bin(), "LibreOffice не установлен")
class Convert(unittest.TestCase):

    def test_pdf_pages_equal_slides(self):
        spec = make_spec()
        fit_deck(spec, THEME)
        pdf = asyncio.run(pptx_to_pdf(render_pptx(spec, watermark=True), timeout=60))
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertEqual(len(re.findall(rb"/Type\s*/Page(?!s)", pdf)), 5)


if __name__ == "__main__":
    unittest.main()
