"""generate_deck целиком на подменном LLM; задача воркера generate_deck_job; маршрутизация в боте."""
import asyncio
import io
import json
import unittest
from types import SimpleNamespace
from unittest import mock

import _helpers  # noqa: F401 — путь к src/ и фиктивные ключи
from _helpers import FakeOpenAI, fake_content, fake_outline

from pptx import Presentation

from config import settings
from core import pipeline
from core.llm import client as llm_client
from core.llm.client import LLMClient
from core.models.request import DeckRequest

TOPIC = "Управление требованиями стейкхолдеров в ИТ-стартапе"


def request(**kw) -> DeckRequest:
    return DeckRequest(input={"topic": kw.pop("topic", TOPIC)}, slides_count=kw.pop("slides_count", 9), **kw)


def run(fake: FakeOpenAI, req: DeckRequest | None = None, **kw):
    async def go():
        return await pipeline.generate_deck("dk_01J0000000000000000000000A", req or request(), client=LLMClient(fake), **kw)
    return asyncio.run(go())


class Pipeline(unittest.TestCase):

    def setUp(self):
        self._model = settings.openai_model
        settings.openai_model = "gpt-6-luna"
        llm_client._SCHEMA_REJECTED.clear()
        self.no_pdf = mock.patch.object(pipeline, "pptx_to_pdf", side_effect=self._fake_pdf)
        self.no_pdf.start()

    def tearDown(self):
        settings.openai_model = self._model
        self.no_pdf.stop()
        llm_client._SCHEMA_REJECTED.clear()

    @staticmethod
    async def _fake_pdf(pptx, **kw):
        return b"%PDF-fake"

    def test_deck_by_topic(self):
        fake = FakeOpenAI()
        res = run(fake)
        spec = res.spec
        self.assertEqual(len(spec.slides), 9)
        self.assertEqual([s.kind for s in spec.slides][::8], ["title", "closing"])
        # G-05: заголовок титула — тема дословно, модель пишет только подзаголовок
        self.assertEqual(spec.slides[0].title, TOPIC)
        self.assertEqual(spec.meta.title, TOPIC)
        self.assertEqual(spec.slides[0].content["subtitle"], "Как свет превращается в сахар")
        self.assertEqual(spec.slides[-1].content["title"], "Спасибо за внимание")
        # сноска у слайдов с числами «по теме» (в fake_content у bullets есть «12 листьях»)
        bullets = [s for s in spec.slides if s.kind == "bullets"]
        self.assertTrue(all(s.footnote == "Оценочные данные — проверьте перед показом" for s in bullets))
        self.assertIsNone(next(s for s in spec.slides if s.kind == "statement").footnote)
        # PPTX читается и в нём 9 слайдов; PDF есть
        self.assertEqual(len(Presentation(io.BytesIO(res.pptx)).slides), 9)
        self.assertEqual(res.pdf, b"%PDF-fake")
        # 1 OUTLINE + 7 CONTENT, стоимость посчитана по этапам
        self.assertEqual(len(fake.calls), 8)
        self.assertEqual(set(res.usage.by_stage), {"outline", "content"})
        self.assertGreater(res.usage.cost_rub, 0)
        for stage in ("ingest", "outline", "select", "content", "fit", "render", "convert", "total"):
            self.assertIn(stage, res.durations_ms)
        self.assertEqual(spec.meta.versions.prompts, open(_helpers.ROOT / "prompts/v2/VERSION").read().strip())

    def test_structured_outputs_request_and_shared_prefix(self):
        fake = FakeOpenAI()
        run(fake)
        formats = [c["response_format"] for c in fake.calls]
        self.assertTrue(all(f["type"] == "json_schema" and f["json_schema"]["strict"] for f in formats))
        self.assertEqual(formats[0]["json_schema"]["name"], "outline")
        systems = {c["messages"][0]["content"] for c in fake.calls[1:]}
        self.assertEqual(len(systems), 1)  # префикс CONTENT одинаков байт в байт
        self.assertEqual(fake.calls[0]["extra_body"]["reasoning_effort"], settings.openai_reasoning_effort)

    def test_fallback_to_json_object_when_provider_rejects_schema(self):
        fake = FakeOpenAI(reject_json_schema=True)
        res = run(fake)
        self.assertEqual(len(res.spec.slides), 9)
        self.assertEqual(fake.calls[0]["response_format"]["type"], "json_schema")
        self.assertTrue(all(c["response_format"] == {"type": "json_object"} for c in fake.calls[1:]))
        self.assertIn("JSON Schema", fake.calls[1]["messages"][0]["content"])

    def test_outline_retry_with_errors_then_ok(self):
        answers = [fake_outline(n=3), fake_outline(n=7)]
        fake = FakeOpenAI(outline=lambda params: answers.pop(0))
        res = run(fake)
        self.assertEqual(len(res.spec.slides), 9)
        retry = fake.calls[1]["messages"]
        self.assertEqual(retry[2]["role"], "assistant")
        self.assertIn("ровно 7", retry[3]["content"])

    def test_outline_failure_is_deck_error(self):
        fake = FakeOpenAI(outline=lambda params: {"oops": True})
        with self.assertRaises(pipeline.DeckError) as ctx:
            run(fake)
        self.assertEqual(ctx.exception.code, "OUTLINE_FAILED")

    def test_content_failure_gives_fallback_slide(self):
        def content(name, user):
            if "СЛАЙД 3 из" in user:
                return {"items": []}  # не проходит minItems дважды
            return fake_content(name, user)
        res = run(FakeOpenAI(content=content))
        s3 = res.spec.slides[2]
        self.assertEqual((s3.kind, s3.variant), ("statement", "statement.big_quote"))
        self.assertEqual(s3.content["body"], "Мысль слайда 2")
        self.assertTrue(any(d.reason == "content_fallback" and d.slide_id == "s03" for d in res.spec.degradations))

    def test_convert_failure_keeps_pptx(self):
        async def broken(pptx, **kw):
            raise pipeline.ConvertError("boom")
        with mock.patch.object(pipeline, "pptx_to_pdf", side_effect=broken):
            res = run(FakeOpenAI())
        self.assertIsNone(res.pdf)
        self.assertTrue(res.pptx)
        self.assertTrue(any(d.stage == "convert" for d in res.spec.degradations))

    def test_watermark_only_in_pdf_copy(self):
        seen = []

        async def capture(pptx, **kw):
            seen.append(pptx)
            return b"%PDF"
        with mock.patch.object(pipeline, "pptx_to_pdf", side_effect=capture):
            res = run(FakeOpenAI(), request(watermark=True))
        texts = lambda data: [sh.text_frame.text for sl in Presentation(io.BytesIO(data)).slides  # noqa: E731
                              for sh in sl.shapes if sh.has_text_frame]
        self.assertNotIn("Fibonacci.", texts(res.pptx))
        self.assertIn("Fibonacci.", texts(seen[0]))

    def test_progress_stages(self):
        stages = []

        async def progress(stage, **kw):
            stages.append((stage, kw.get("done")))
        run(FakeOpenAI(), progress=progress)
        names = [s for s, _ in stages]
        self.assertEqual(names[:3], ["ingest", "outline", "content"])
        self.assertEqual(names[-2:], ["render", "convert"])
        self.assertIn(("content", 7), stages)


# ── Воркер ───────────────────────────────────────────────────────────────────

class FakeBot:
    def __init__(self):
        self.edits, self.albums, self.documents, self.messages, self.deleted = [], [], [], [], []

    async def edit_message_text(self, text, chat_id=None, message_id=None, parse_mode=None):
        self.edits.append(text)

    async def send_media_group(self, chat_id, media):
        self.albums.append(media)

    async def send_document(self, chat_id, document=None, caption=None, parse_mode=None):
        self.documents.append((document, caption))

    async def send_message(self, chat_id, text, reply_markup=None):
        self.messages.append((text, reply_markup))

    async def delete_message(self, chat_id=None, message_id=None):
        self.deleted.append(message_id)


class WorkerJob(unittest.TestCase):

    def _req(self):
        return request(client={"user_id": 1, "chat_id": 1, "status_message_id": 10}, watermark=True).model_dump(mode="json")

    def test_job_delivers_album_and_counts(self):
        import worker
        from core.models.deck import DeckSpec

        bot = FakeBot()
        fake = FakeOpenAI()

        async def gen(deck_id, req, progress):
            await progress("outline")
            await progress("content", done=1, total=7)
            with mock.patch.object(pipeline, "pptx_to_pdf", side_effect=Pipeline._fake_pdf):
                return await pipeline.generate_deck(deck_id, req, progress, client=LLMClient(fake))

        finished = {}

        async def finish(deck_id, **kw):
            finished.update(kw)

        with mock.patch.object(worker, "generate_deck", side_effect=gen), \
                mock.patch.object(worker.deck_store, "finish", side_effect=finish):
            out = asyncio.run(worker.generate_deck_job({"bot": bot}, "dk_01J0000000000000000000000A",
                                                        request=self._req()))
        self.assertEqual(out["status"], "ok")
        self.assertEqual(len(bot.albums), 1)
        names = [m.media.filename for m in bot.albums[0]]
        self.assertTrue(names[0].endswith(".pptx") and names[1].endswith(".pdf"))
        self.assertIn("Бесплатная версия", bot.albums[0][-1].caption)
        self.assertIn("🧭 Составляю план", bot.edits[0])
        self.assertEqual(bot.deleted, [10])
        self.assertEqual(finished["status"], "done")
        DeckSpec.model_validate(finished["spec"])
        self.assertIn("outline", finished["usage"])

    def test_job_reports_outline_failure(self):
        import worker
        bot = FakeBot()
        failed = []

        async def gen(deck_id, req, progress):
            raise pipeline.DeckError("OUTLINE_FAILED")

        async def fail(deck_id, code, **kw):
            failed.append(code)
        with mock.patch.object(worker, "generate_deck", side_effect=gen), \
                mock.patch.object(worker.deck_store, "fail", side_effect=fail):
            out = asyncio.run(worker.generate_deck_job({"bot": bot}, "dk_x", request=self._req()))
        self.assertEqual(out["error_code"], "OUTLINE_FAILED")
        self.assertEqual(failed, ["OUTLINE_FAILED"])
        self.assertIn("Не удалось составить план", bot.edits[-1])
        self.assertEqual(bot.albums, [])


# ── Бот: маршрутизация между движками (D-038) ───────────────────────────────

class Routing(unittest.TestCase):

    BASE = {"topic": "Как работает фотосинтез", "audience": "students", "language": "ru", "color_scheme": "dark",
            "slide_count": 9, "source_mode": "strict"}

    def test_engine_for(self):
        import dialog
        self.assertEqual(dialog.engine_for({**self.BASE, "presentation_type": "doklad"}), "new")
        self.assertEqual(dialog.engine_for({**self.BASE, "presentation_type": "doklad", "source_type": "text"}), "old")
        self.assertEqual(dialog.engine_for({**self.BASE, "presentation_type": "doklad", "source_type": "document"}), "old")
        self.assertEqual(dialog.engine_for({**self.BASE, "presentation_type": "pitch_deck"}), "old")

    def test_summary_for_new_engine(self):
        import dialog
        data = {**self.BASE, "presentation_type": "doklad"}
        text = dialog.summary_text(data, "free")
        self.assertIn("Графит светлая", text)
        self.assertIn("PPTX + PDF", text)
        kb = [b.callback_data for row in dialog.summary_keyboard(data).inline_keyboard for b in row]
        self.assertNotIn("sum:design", kb)
        self.assertIn("sum:slides", kb)
        old = {**data, "source_type": "text", "raw_text": "x" * 300}
        self.assertIn("sum:design", [b.callback_data for row in dialog.summary_keyboard(old).inline_keyboard for b in row])

    def test_pitch_deck_has_no_slide_count_button(self):
        import dialog
        kb = [b.callback_data for row in dialog.summary_keyboard({**self.BASE, "presentation_type": "pitch_deck"})
              .inline_keyboard for b in row]
        self.assertNotIn("sum:slides", kb)

    def test_theme_mapping(self):
        import dialog
        self.assertEqual(dialog.theme_for("light"), "graphite_light")
        self.assertEqual(dialog.theme_for("dark"), "graphite_light")    # тёмная ещё не включена
        self.assertEqual(dialog.theme_for("dark", ["graphite_light", "graphite_dark"]), "graphite_dark")
        self.assertEqual(dialog.theme_for("forest"), "graphite_light")

    def test_topic_threshold_200(self):
        import dialog
        self.assertEqual(dialog.TOPIC_MAX_CHARS, 200)

    def test_new_engine_enqueues_deck_job(self):
        import main
        from test_dialog import Chat, FakeMessage

        jobs, rows = [], []

        class Pool:
            async def enqueue_job(self, name, *args, **kw):
                jobs.append((name, args, kw))

        async def create(deck_id, payload, user_id, parent_deck_id=None):
            rows.append((deck_id, payload, user_id))
            return True

        chat = Chat()
        with mock.patch.object(main, "arq_pool", Pool()), mock.patch.object(main.deck_store, "create_deck", side_effect=create):
            asyncio.run(main.generate_and_send(FakeMessage(chat), {**self.BASE, "presentation_type": "doklad"},
                                               watermark=True))
        name, args, kw = jobs[0]
        self.assertEqual(name, "generate_deck_job")
        deck_id = args[0]
        self.assertRegex(deck_id, r"^dk_")
        self.assertIsNone(kw["request"])       # параметры — в строке decks, в очереди только deck_id
        self.assertEqual(kw["_job_id"], deck_id)
        payload = rows[0][1]
        req = DeckRequest.model_validate(payload)
        self.assertEqual((req.input.topic, req.theme_id, req.slides_count, req.watermark, req.audience),
                         ("Как работает фотосинтез", "graphite_light", 9, True, "students"))
        self.assertEqual(req.client.chat_id, chat.id)
        self.assertEqual(req.client.status_message_id, chat.sent[0].message_id)
        self.assertIn(deck_id, chat.sent[0].body)

    def test_without_db_request_goes_into_job(self):
        import main
        from test_dialog import Chat, FakeMessage
        jobs = []

        class Pool:
            async def enqueue_job(self, name, *args, **kw):
                jobs.append(kw)

        async def create(*a, **kw):
            return False
        with mock.patch.object(main, "arq_pool", Pool()), mock.patch.object(main.deck_store, "create_deck", side_effect=create):
            asyncio.run(main.generate_and_send(FakeMessage(Chat()), {**self.BASE, "presentation_type": "doklad"}))
        self.assertEqual(jobs[0]["request"]["input"]["topic"], "Как работает фотосинтез")


if __name__ == "__main__":
    unittest.main()
