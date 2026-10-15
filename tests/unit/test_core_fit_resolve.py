"""Текст никогда не режется (сессия 5, D-071…D-074): перенос как у LibreOffice, типографика чисел,
невлезание по порядку «вариант → сокращение → два слайда», финальная проверка готового PPTX,
мусор генерации, заполненность зоны контента."""
import asyncio
import io
import unittest

import _helpers  # noqa: F401 — путь к src/ и фиктивные ключи

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Emu, Pt

from core.checks.layout import check_pptx, fill_share
from core.checks.text import Known, problems, remove_fragments
from core.content.fill import count_remarks, length_remarks, BulletsContent, BulletItem
from core.fitting.fitter import fit_deck, layout_slide
from core.fitting.metrics import NBSP, WJ, overlong_words, text_width, wrap
from core.fitting.resolve import FitError, regroup, resolve_deck, split_slide
from core.models.deck import DeckSpec, Meta, Slide
from core.models.ids import new_deck_id
from core.models.layout import load_layouts
from core.models.theme import load_theme
from core.render.pptx.renderer import render_pptx
from core.text.typography import plain, typeset

L = load_layouts()
THEME = load_theme("business_slate")
SURGERY_POINTS = [
    "Универсальные платформы требуют интеграции с оборудованием операционной и обучения хирургов",
    "Специализированное решение работает с существующими эндоскопами без замены стойки",
    "Внедрение занимает 4–8 мес. и требует регистрационного удостоверения Росздравнадзора",
    "Точность 92–98 % подтверждена на видео лапароскопических операций трёх клиник",
]


def spec_with(*slides: Slide, title: str = "AI-навигация для эндоскопической хирургии") -> DeckSpec:
    body = [Slide(id="s01", index=1, kind="title", variant="title.cover_center", title=title,
                  content={"subtitle": "Как компьютерное зрение помогает хирургу", "author": None})]
    body += list(slides)
    body.append(Slide(id="s99", index=99, kind="closing", variant="closing.thanks_center", title="Спасибо за внимание",
                      content={"title": "Спасибо за внимание", "body": None, "author": None}))
    while len(body) < 4:
        body.insert(1, Slide(id="s98", index=98, kind="statement", variant="statement.big_quote",
                             title="Короткое утверждение", content={"body": None}))
    for i, s in enumerate(body, start=1):
        s.index, s.id = i, f"s{i:02d}"
    return DeckSpec(id=new_deck_id(), meta=Meta(title=title, language="ru", presentation_type="doklad",
                                                audience="general", mode="topic", theme_id="business_slate", seed=1),
                    slides=body)


def all_text(spec: DeckSpec) -> str:
    parts = []
    for s in spec.slides:
        parts.append(s.title)

        def walk(v):
            if isinstance(v, str):
                parts.append(v)
            elif isinstance(v, dict):
                for x in v.values():
                    walk(x)
            elif isinstance(v, list):
                for x in v:
                    walk(x)
        walk(s.content)
    return "\n".join(parts)


class Wrap(unittest.TestCase):

    def test_double_space_after_bullet_is_measured(self):
        """Причина наложения в comparison: «•  » мерилось с одним пробелом — строка у LibreOffice
        шире, чем считал fitter, и пункт уходил на лишнюю строку поверх следующего."""
        one, two = text_width("• текст", 18, False), text_width("•  текст", 18, False)
        self.assertGreater(two, one)
        line = wrap("•  слово " * 3, 10_000, 18, False)[0]
        self.assertTrue(line.startswith("•  слово"))

    def test_nbsp_and_word_joiner_do_not_break(self):
        text = f"время 60{WJ}–{WJ}90{NBSP}секунд и точность 92{WJ}–{WJ}98{NBSP}%"
        for width in (120, 160, 200, 260):
            lines = wrap(text, width, 18, False)
            joined = [line for line in lines if "60" in line or "90" in line]
            self.assertTrue(any(f"60{WJ}–{WJ}90{NBSP}секунд" in line for line in joined) or
                            overlong_words(text, width, 18, False), (width, lines))
            self.assertFalse(any(line.endswith("–") for line in lines), lines)
        self.assertEqual(text_width(WJ, 18, False), 0)

    def test_hyphen_is_a_break_opportunity(self):
        self.assertEqual(wrap("в ИТ-стартапах и научно-исследовательских", 330, 18, False)[-2:],
                         ["научно-", "исследовательских"])

    def test_overlong_word(self):
        self.assertEqual(overlong_words("Непротиворечивость данных", 150, 18, False), ["Непротиворечивость"])
        self.assertEqual(overlong_words("Непротиворечивость данных", 400, 18, False), [])


class Typography(unittest.TestCase):

    def test_ranges_units_percent(self):
        self.assertEqual(typeset("за 60–90 секунд", "ru"), f"за 60{WJ}–{WJ}90{NBSP}секунд")
        self.assertEqual(typeset("точность 92–98 %", "ru"), f"точность 92{WJ}–{WJ}98{NBSP}%")
        self.assertEqual(typeset("4–8 мес.: интеграция", "ru"), f"4{WJ}–{WJ}8{NBSP}мес.: интеграция")
        self.assertEqual(typeset("1 250 000 пациентов", "ru"), f"1{NBSP}250{NBSP}000{NBSP}пациентов")

    def test_decimal_comma_by_language(self):
        self.assertEqual(plain(typeset("цель — до 98.79 %", "ru")), "цель — до 98,79 %")
        self.assertEqual(plain(typeset("goal 98.79 %", "en")), "goal 98.79 %")
        for keep in ("Python 3.11", "дата 12.05.2026", "GPT-4.5", "v2.0"):
            self.assertEqual(plain(typeset(keep, "ru")), keep)

    def test_phone_is_not_a_range(self):
        self.assertEqual(typeset("+7-999-123-45", "ru"), "+7-999-123-45")


class TextGarbage(unittest.TestCase):
    """Найденное владельцем и golden PR #83: иероглиф, хангыль, «�», обрывок «Iн», «liver», склейка."""

    known = Known.build("Boston Scientific и Medtronic, лицензия Apache 2.0 или MIT, end-to-end")

    def codes(self, text):
        return [(p.code, p.fragment) for p in problems(text, "ru", self.known)]

    def test_detects(self):
        self.assertEqual(self.codes("модели до 子"), [("script", "子")])
        self.assertEqual(self.codes("люди и за다"), [("script", "다")])
        self.assertEqual(self.codes("текст с �"), [("script", "�")])
        self.assertIn(("mixed", "Iн"), self.codes("Iн тегрaция"))
        self.assertEqual(self.codes("рост earnings и liver"), [("foreign", "earnings"), ("foreign", "liver")])
        self.assertEqual([(p.code, p.fix) for p in problems("BostonScientific лидирует", "ru", self.known)],
                         [("glued", "Boston Scientific")])

    def test_allowed(self):
        for ok in ("end-to-end прогон, AI, API, KPI, Python, SaaS", "лицензия apache 2.0 или mit",
                   "AI-навигация для IT-компаний, e-mail",
                   "Boston Scientific и Medtronic", "оценка α и β", "₽, №, °C, ×"):
            self.assertEqual(self.codes(ok), [], ok)
        # «_» в имени продукта — не чужой знак (golden PR #85, G-06: «sechenov.ai_nephro»)
        self.assertNotIn("script", [c for c, _ in self.codes("прецедент Sechenov.AI_Nephro получил РУ")])

    def test_english_deck_flags_cyrillic(self):
        self.assertEqual([p.code for p in problems("growth of выручка", "en", Known())], ["foreign"])

    def test_remove_fragments(self):
        text = "модели — с лицензией MIT, до 子"
        self.assertEqual(remove_fragments(text, problems(text, "ru", self.known)), "модели — с лицензией MIT, до")


class ContentChecks(unittest.TestCase):

    def test_length_and_count_are_remarks_not_schema(self):
        layout = L["bullets.cards_grid"]
        long = BulletsContent(items=[BulletItem(heading="Пункт", text="слово " * 80, icon="bulb")] * 7)
        remarks = length_remarks("bullets", long, layout, 6)
        self.assertTrue(any("сократи до" in r for r in remarks))
        self.assertTrue(count_remarks("bullets", long, layout, 6))
        cut = BulletsContent(items=[BulletItem(heading="Пункт", text="Оборвано…", icon="bulb")] * 3)
        self.assertTrue(any("многоточие" in r for r in length_remarks("bullets", cut, layout, 3)))


class Resolve(unittest.TestCase):

    def run_resolve(self, spec, shorten=None):
        overflows = fit_deck(spec, THEME)
        asyncio.run(resolve_deck(spec, THEME, overflows, shorten=shorten))
        return overflows

    def test_no_ellipsis_anywhere(self):
        long = ("Алгоритм распознаёт анатомические структуры и подсказывает хирургу безопасную траекторию "
                "инструмента во время лапароскопических операций ") * 3
        slide = Slide(id="s02", index=2, kind="bullets", variant="bullets.cards_grid", title=long[:180],
                      content={"items": [{"heading": "Интеграция", "text": long, "icon": "bulb"}] * 4})
        spec = spec_with(slide)
        spec.slides[0].content["subtitle"] = long

        async def shorten(sl, targets):
            return {t.id: " ".join(t.text.split()[:4]) for t in targets}
        self.run_resolve(spec, shorten)
        self.assertNotIn("…", all_text(spec))
        self.assertFalse(fit_deck(spec, THEME))                      # всё влезло

    def test_variant_before_shortening(self):
        """Длинное слово в узкой карточке 3×2 → вариант, где влезает, без LLM."""
        items = [{"heading": "Непротиворечивостьтребований", "text": "Коротко.", "icon": "bulb"} for _ in range(4)]
        slide = Slide(id="s02", index=2, kind="bullets", variant="bullets.numbered_columns", title="Требования",
                      content={"items": items})
        spec = spec_with(slide)
        calls = []

        async def shorten(sl, targets):
            calls.append(targets)
            return None
        overflows = self.run_resolve(spec, shorten)
        self.assertTrue(overflows)
        self.assertEqual(calls, [])
        self.assertNotEqual(spec.slides[1].variant, "bullets.numbered_columns")
        self.assertTrue(any(d.reason.startswith("variant:") for d in spec.degradations))

    def test_shorten_then_split(self):
        text = "Хирург видит границы протоков и сосудов прямо на экране эндоскопа во время операции. " * 5
        items = [{"text": text} for _ in range(4)]
        slide = Slide(id="s02", index=2, kind="conclusion", variant="conclusion.numbered_takeaways",
                      title="Выводы", content={"items": items})
        spec = spec_with(slide)

        async def no_luck(sl, targets):
            return {t.id: t.text for t in targets}                 # «модель» не сократила
        self.run_resolve(spec, no_luck)
        conclusions = [s for s in spec.slides if s.kind == "conclusion"]
        self.assertEqual(len(conclusions), 2)                        # два слайда, пункты не потеряны
        self.assertEqual(sum(len(s.content["items"]) for s in conclusions), 4)
        self.assertIn("(продолжение)", conclusions[1].title)
        self.assertEqual([s.index for s in spec.slides], list(range(1, len(spec.slides) + 1)))

    def test_topic_is_never_shortened(self):
        topic = "Итоги пилота «Навигатор»: что изменилось для сотрудников за полгода работы и что делать дальше"
        spec = spec_with(title=topic)
        spec.slides[0].variant = "title.cover_split_image"
        spec.slides[0].image = "img_1"

        async def shorten(sl, targets):
            self.fail("тема пользователя не сокращается")
        self.run_resolve(spec, shorten)
        self.assertEqual(spec.meta.title, topic)

    def test_fit_error_when_nothing_helps(self):
        word = "Непротиворечивость" * 6
        slide = Slide(id="s02", index=2, kind="statement", variant="statement.big_quote", title=word,
                      content={"body": None})
        spec = spec_with(slide)
        with self.assertRaises(FitError):
            self.run_resolve(spec, None)

    def test_split_and_regroup_keep_items(self):
        slide = Slide(id="s02", index=2, kind="bullets", variant="bullets.icon_list", title="Пункты",
                      content={"items": [{"heading": f"П{i}", "text": "Текст.", "icon": "bulb"} for i in range(8)]})
        parts = regroup(slide, "ru")
        self.assertEqual(sum(len(p.content["items"]) for p in parts), 8)
        self.assertTrue(all(len(p.content["items"]) in L[p.variant].arrangements for p in parts))
        self.assertIsNone(split_slide(Slide(id="s02", index=2, kind="statement", variant="statement.big_quote",
                                            title="x", content={}), "ru"))


class FinalCheck(unittest.TestCase):
    """Финальная проверка готового PPTX ловит то, что владелец находил глазами."""

    @staticmethod
    def pptx_with(*boxes) -> bytes:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12_192_000), Emu(6_858_000)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for name, (x, y, w, h), text, pt in boxes:
            sh = slide.shapes.add_textbox(Emu(x * 6350), Emu(y * 6350), Emu(w * 6350), Emu(h * 6350))
            sh.name = name
            tf = sh.text_frame
            tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.line_spacing = Pt(pt * 1.2)
            run = p.add_run()
            run.text, run.font.size, run.font.name = text, Pt(pt), "Arial"
        out = io.BytesIO()
        prs.save(out)
        return out.getvalue()

    def codes(self, pptx):
        return sorted({(i.shape, i.code) for i in check_pptx(pptx)})

    def test_overflow_overlap_word_ellipsis(self):
        point = "•  " + SURGERY_POINTS[0]
        pptx = self.pptx_with(("left_point1", (120, 400, 400, 43), point, 18),
                              ("left_point2", (120, 450, 400, 43), "•  Следующий пункт", 18),
                              ("heading#1", (120, 600, 150, 43), "Непротиворечивость", 18),
                              ("text#1", (120, 700, 800, 43), "Время подготовки снимка сократилось с…", 18))
        codes = self.codes(pptx)
        self.assertIn(("left_point1", "overflow"), codes)
        self.assertIn(("left_point2", "overlap"), codes)
        self.assertIn(("heading#1", "word"), codes)
        self.assertIn(("text#1", "ellipsis"), codes)

    def test_comparison_slide_11_has_no_overlap(self):
        """Хирургия, слайд 11: «универсальные» поверх следующего пункта — во всех темах чисто."""
        for theme_id in ("business_slate", "ember_dark", "sunny_cream", "mint_coral"):
            for variant in ("comparison.two_columns", "comparison.pros_cons"):
                slide = Slide(id="s02", index=2, kind="comparison", variant=variant,
                              title="Универсальные платформы против специализированного решения",
                              content={"left": {"header": "Универсальные платформы", "points": SURGERY_POINTS},
                                       "right": {"header": "Специализированное решение", "points": SURGERY_POINTS[::-1]},
                                       "polarity": "pros_cons"})
                spec = spec_with(slide)
                spec.meta.theme_id = theme_id
                for s in spec.slides:
                    s.content = {k: (typeset(v) if isinstance(v, str) else v) for k, v in s.content.items()}
                overflows = fit_deck(spec, load_theme(theme_id))
                asyncio.run(resolve_deck(spec, load_theme(theme_id), overflows))
                self.assertEqual(check_pptx(render_pptx(spec)), [], (theme_id, variant))


class Fill(unittest.TestCase):

    def test_sparse_cards_stretch_to_40_percent(self):
        slide = Slide(id="s02", index=2, kind="bullets", variant="bullets.cards_grid", title="Три пункта",
                      content={"items": [{"heading": "Коротко", "text": "Одна фраза.", "icon": "bulb"}] * 3})
        spec = spec_with(slide)
        layout_slide(spec.slides[1], spec, L["bullets.cards_grid"], THEME)
        prs = Presentation(io.BytesIO(render_pptx(spec)))
        self.assertGreaterEqual(fill_share(prs.slides[1]), 40.0)


if __name__ == "__main__":
    unittest.main()
