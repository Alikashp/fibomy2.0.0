"""Материал в новом движке: INGEST, данные диаграмм, проверка чисел, разнообразие плана, колода целиком."""
import asyncio
import io
import unittest
import zipfile
from unittest import mock

import _helpers  # noqa: F401 — путь к src/ и фиктивные ключи
from _helpers import FIXTURES, FakeOpenAI, fake_content, fake_outline

from pptx import Presentation

from config import settings
from core import pipeline
from core.checks.facts import slide_errors, strip_unknown
from core.ingest import IngestError, ingest, parse
from core.ingest.numbers import find_numbers, format_number, parse_cell
from core.ingest.render import digest_text
from core.ingest.tables import build_dataset, text_series
from core.llm import client as llm_client
from core.llm.client import LLMClient
from core.models.deck import Slide
from core.models.outline import OutlineResponse, PlannedDataset, PlannedSlide
from core.models.request import DeckRequest
from core.planning.datasets import DatasetError, legend, resolve
from core.planning.outline import diversity_remarks, normalize_plan, plan_remarks, available_kinds

NAV = (FIXTURES / "navigator_pilot.txt").read_bytes()
DOCX = (FIXTURES / "test_prodazhi_Q3.docx").read_bytes()


class Numbers(unittest.TestCase):

    def test_parse_cell(self):
        cases = {"47,3": (47.3, False), "1 150": (1150, False), "1 150": (1150, False), "−6": (-6, False),
                 "+41%": (41, True), "65": (65, False), "новый канал": (None, False), "4,5 млн": (None, False)}
        for raw, expected in cases.items():
            self.assertEqual(parse_cell(raw), expected, raw)

    def test_find_numbers(self):
        values = [v for v, _, _ in find_numbers("Рост с 310 до 2 050, доля 46%, бюджет 4,5 млн, 2026–2027")]
        self.assertEqual(values, [310, 2050, 46, 4.5, 2026, 2027])

    def test_format_number(self):
        self.assertEqual(format_number(2050), "2 050")
        self.assertEqual(format_number(47.3), "47,3")
        self.assertEqual(format_number(-6), "−6")
        self.assertEqual(format_number(4.5, "en"), "4.5")


class Ingest(unittest.TestCase):

    def test_docx_tables_become_datasets(self):
        d = parse(DOCX, "docx", "test_prodazhi_Q3.docx")
        self.assertEqual(len(d.datasets), 3)
        months, channels, cities = d.datasets
        self.assertEqual((months.axis, months.title), ("time", "Таблица 1. Помесячные показатели"))
        self.assertTrue(months.rows[-1].is_total)
        self.assertEqual((channels.columns[1].name, channels.columns[1].unit), ("Выручка", "млн руб."))
        self.assertEqual(channels.columns[2].type, "percent")
        self.assertEqual(channels.sums["c3"], 100.0)
        self.assertEqual(channels.rows[3].cells["c4"].value, -6)
        self.assertIsNone(channels.rows[2].cells["c4"].value)       # «новый канал»
        self.assertEqual(cities.axis, "category")
        self.assertEqual(d.fragments[0].kind, "heading")

    def test_text_series_from_navigator(self):
        d = parse(NAV, "text")
        users, shares = d.datasets
        self.assertEqual([r.cells["c2"].value for r in users.rows], [310, 720, 1150, 1540, 1820, 2050])
        self.assertEqual(users.axis, "time")
        self.assertEqual(shares.sums["c2"], 100.0)
        self.assertEqual([r.label for r in shares.rows], ["документы", "переписки", "люди и контакты", "задачи"])
        self.assertIsNone(text_series("Среднее время поиска: до пилота 4,2 минуты, после — 1,1 минуты."))

    def test_pdf_text(self):
        d = parse((FIXTURES / "test_hakaton_Q3.pdf").read_bytes(), "pdf", "x.pdf")
        self.assertGreater(d.source.chars_used, 8000)
        self.assertEqual(d.source.pages, 7)

    def test_pptx_table_and_chart(self):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Отгрузки складов за год: север растёт быстрее юга"
        table = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
        for r, row in enumerate([["Склад", "Тонн"], ["Север", "120"], ["Юг", "95"]]):
            for c, v in enumerate(row):
                table.cell(r, c).text = v
        data = CategoryChartData()
        data.categories = ["I кв.", "II кв.", "III кв."]
        data.add_series("Тонн", (10, 12, 15))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(3), data)
        buf = io.BytesIO()
        prs.save(buf)
        d = parse(buf.getvalue(), "pptx", "x.pptx")
        origins = [ds.origin for ds in d.datasets]
        self.assertEqual(origins, ["pptx_table", "pptx_chart"])
        self.assertEqual(d.datasets[1].axis, "time")
        self.assertEqual([r.cells["c2"].value for r in d.datasets[1].rows], [10, 12, 15])

    def test_errors(self):
        with self.assertRaises(IngestError) as ctx:
            parse(b"not a zip", "docx")
        self.assertEqual(ctx.exception.code, "BAD_FILE")
        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("big.xml", b"0" * (210 * 1024 * 1024))
        with self.assertRaises(IngestError) as ctx:
            parse(bomb.getvalue(), "docx")
        self.assertEqual(ctx.exception.code, "BAD_FILE")
        with self.assertRaises(IngestError) as ctx:
            parse(b"   ", "text")
        self.assertEqual(ctx.exception.code, "BAD_FILE")

    def test_scan_pdf(self):
        pdf = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
               b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF")
        with self.assertRaises(IngestError) as ctx:
            parse(pdf, "pdf")
        self.assertEqual(ctx.exception.code, "SCAN_WITHOUT_TEXT")

    def test_truncation(self):
        text = ("Абзац про склад и отгрузки. " * 20 + "\n\n") * 100
        d = parse(text.encode(), "text")
        self.assertTrue(d.source.truncated)
        self.assertLessEqual(d.source.chars_used, 40_000)
        self.assertTrue(any(w.startswith("truncated:") for w in d.warnings))

    def test_ingest_request(self):
        req = DeckRequest(input={"topic": "Пилот", "material": {"kind": "text", "ref": "redis:x"}})
        d = asyncio.run(ingest(req, NAV))
        self.assertEqual((d.mode, d.topic), ("material", "Пилот"))
        self.assertIn("[ds2] Структура запросов за июнь", digest_text(d))
        self.assertIn("<source>", digest_text(d))
        doc = DeckRequest(input={"topic": "Тема", "material": {"kind": "document", "ref": "s3:x",
                                                            "mime": "application/zip", "name": "a.rar"}})
        with self.assertRaises(IngestError) as ctx:
            asyncio.run(ingest(doc, b"x"))
        self.assertEqual(ctx.exception.code, "FILE_UNSUPPORTED")


def planned(kind, ds, cols=("c2",), rows=None):
    return PlannedSlide(kind=kind, title="Т", key_message="К",
                        dataset=PlannedDataset(id=ds, label_column="c1", value_columns=list(cols), rows=rows))


class ChartData(unittest.TestCase):

    def setUp(self):
        self.nav = parse(NAV, "text")
        self.docx = parse(DOCX, "docx")

    def test_series_and_share(self):
        data = resolve(planned("chart_series", "ds1"), self.nav, "chart_series")
        self.assertEqual(data["series"][0]["values"], [310, 720, 1150, 1540, 1820, 2050])
        self.assertEqual(data["axis"], "time")
        share = resolve(planned("chart_share", "ds2"), self.nav, "chart_share")
        self.assertEqual(share["share_sum"], 100.0)
        self.assertEqual([x["value"] for x in legend(share)], ["46%", "27%", "18%", "9%"])

    def test_total_row_excluded_and_decimals(self):
        data = resolve(planned("chart_series", "ds1"), self.docx, "chart_series")
        self.assertEqual([c["label"] for c in data["categories"]], ["Июль", "Август", "Сентябрь"])
        self.assertEqual(data["decimals"], 1)

    def test_share_from_values_and_bad_sum(self):
        data = resolve(planned("chart_share", "ds3", ("c3",)), self.docx, "chart_share")  # выручка городов
        self.assertTrue(data["share_from_values"])
        self.assertAlmostEqual(sum(data["series"][0]["values"]), 100, delta=0.5)
        with self.assertRaises(DatasetError):  # «Выполнение плана, %» — сумма 401
            resolve(planned("chart_share", "ds3", ("c4",)), self.docx, "chart_share")
        with self.assertRaises(DatasetError):
            resolve(planned("chart_series", "ds9"), self.docx, "chart_series")

    def test_bad_share_falls_back_to_columns(self):
        resp = OutlineResponse.model_validate(fake_outline(
            n=3, kinds=["chart_share", "bullets", "conclusion"],
            datasets={0: {"id": "ds3", "label_column": "c1", "value_columns": ["c4"], "rows": None}}))
        req = DeckRequest(input={"topic": "Тема", "material": {"kind": "text", "ref": "x"}}, slides_count=5)
        reasons = []
        slides, datas = normalize_plan(resp, req, 3, available_kinds(req), reasons.append, self.docx)
        self.assertEqual(slides[0].kind, "chart_series")
        self.assertEqual(datas[0]["series"][0]["values"], [112, 104, 87, 98])
        self.assertTrue(any(r.startswith("chart_share->chart_series") for r in reasons))


class Facts(unittest.TestCase):

    def setUp(self):
        self.digest = parse(NAV, "text")

    def slide(self, kind, content, title="Пилот охватил 2 400 сотрудников"):
        return Slide(id="s02", index=2, kind=kind, variant="metrics.kpi_cards", title=title, content=content)

    def test_known_numbers_pass(self):
        s = self.slide("metrics", {"items": [
            {"value": "2 400", "label": "сотрудников", "source_ref": "f7"},
            {"value": "8 000", "label": "цель 2027", "source_ref": None},
            {"value": "46%", "label": "документы", "source_ref": "ds2:r1c2"}], "body": "Экономия 40 минут в день."})
        self.assertEqual(slide_errors(s, self.digest), [])

    def test_invented_and_mismatched(self):
        s = self.slide("metrics", {"items": [
            {"value": "в 6,6 раза", "label": "рост пользователей", "source_ref": None},
            {"value": "27%", "label": "документы", "source_ref": "ds2:r1c2"},
            {"value": "74%", "label": "быстрее поиск", "source_ref": None}], "body": None})
        errors = slide_errors(s, self.digest)
        self.assertTrue(any("6,6" in e for e in errors))
        self.assertTrue(any("74%" in e for e in errors))
        self.assertTrue(any("не совпадает с ячейкой ds2:r1c2" in e for e in errors))
        strip_unknown(s, self.digest)
        self.assertEqual([i["value"] for i in s.content["items"]], ["27%"])

    def test_small_integers_ignored_and_sentences_stripped(self):
        s = Slide(id="s03", index=3, kind="bullets", variant="bullets.cards_grid", title="Пять систем",
                  content={"items": [{"heading": "Поиск", "text": "Ищет по 5 системам. Ускоряет работу на 85%.",
                                      "icon": "search"}]})
        self.assertEqual(len(slide_errors(s, self.digest)), 1)
        strip_unknown(s, self.digest)
        self.assertEqual(s.content["items"][0]["text"], "Ищет по 5 системам.")


class Diversity(unittest.TestCase):

    def test_rules(self):
        self.assertEqual(diversity_remarks(["statement", "bullets", "process", "comparison", "metrics",
                                            "bullets", "conclusion"]), [])
        remarks = diversity_remarks(["bullets", "bullets", "statement", "bullets", "bullets", "process", "conclusion"])
        self.assertTrue(any("bullets на 4 слайдах из 7 — не больше 2" in r for r in remarks))
        self.assertTrue(any("слайды 1 и 2 подряд" in r for r in remarks))
        self.assertTrue(any("statement на 4" in r for r in diversity_remarks(["statement", "bullets"] * 3 + ["statement"])))

    def test_plan_remarks_include_dataset_errors(self):
        nav = parse(NAV, "text")
        req = DeckRequest(input={"topic": "Тема", "material": {"kind": "text", "ref": "x"}})
        resp = OutlineResponse.model_validate(fake_outline(
            n=3, kinds=["chart_series", "chart_series", "conclusion"],
            datasets={0: {"id": "ds1", "label_column": "c1", "value_columns": ["c2"], "rows": None},
                      1: {"id": "ds1", "label_column": "c1", "value_columns": ["c2"], "rows": None}}))
        remarks = plan_remarks(resp, req, nav)
        self.assertTrue(any("одна колонка ds1:c2 на двух диаграммах" in r for r in remarks))
        self.assertTrue(any("подряд" in r for r in remarks))

    def test_soft_remarks_retry_then_accept(self):
        """Замечание о разнообразии → повтор; на последней попытке план принимается."""
        settings_model = settings.openai_model
        settings.openai_model = "gpt-6-luna"
        try:
            monotone = fake_outline(kinds=["bullets"] * 6 + ["conclusion"])
            fake = FakeOpenAI(outline=monotone)
            req = DeckRequest(input={"topic": "Как устроен город"}, slides_count=9)
            with mock.patch.object(pipeline, "pptx_to_pdf", side_effect=_pdf):
                res = asyncio.run(pipeline.generate_deck("dk_01J0000000000000000000000A", req,
                                                         client=LLMClient(fake)))
            outline_calls = [c for c in fake.calls if c["response_format"].get("json_schema", {}).get("name") == "outline"]
            self.assertEqual(len(outline_calls), 2)
            self.assertIn("bullets на 6 слайдах", outline_calls[1]["messages"][-1]["content"])
            self.assertTrue(any(d.reason.startswith("plan_diversity") for d in res.spec.degradations))
        finally:
            settings.openai_model = settings_model


async def _pdf(pptx, **kw):
    return b"%PDF"


NAV_PLAN = fake_outline(
    n=7, genre="report", subtitle="Отчёт руководству",
    kinds=["statement", "metrics", "chart_series", "chart_share", "comparison", "process", "statement"],
    asks=[{"text": "Утвердить бюджет на масштабирование — 4,5 млн руб.", "refs": ["f25"]}],
    datasets={2: {"id": "ds1", "label_column": "c1", "value_columns": ["c2"], "rows": None},
              3: {"id": "ds2", "label_column": "c1", "value_columns": ["c2"], "rows": None}})
NAV_PLAN["deck"]["slides"][-1]["role"] = "ask"
NAV_PLAN["deck"]["slides"][-1]["title"] = "Нужно решение: утвердить бюджет на масштабирование — 4,5 млн руб."


def nav_content(name, user):
    if name == "slide_metrics":
        if "ОШИБКИ ПРОВЕРКИ" in user:
            return {"items": [{"value": "2 400", "unit": "сотрудников", "label": "в пилоте", "status": "fact",
                               "source_ref": "f7"},
                              {"value": "8 000", "unit": None, "label": "активных пользователей в месяц",
                               "status": "plan", "source_ref": "f24"}], "body": None}
        return {"items": [{"value": "2 400", "unit": "сотрудников", "label": "в пилоте", "status": "fact",
                           "source_ref": "f7"},
                          {"value": "в 6,6 раза", "unit": None, "label": "рост пользователей", "status": "fact",
                           "source_ref": None}], "body": None}
    if name == "slide_chart_share":
        return {"insight": "Почти половина запросов — документы.", "value_axis_title": None,
                "category_labels": [{"row": "r3", "label": "люди"}]}
    return fake_content(name, user)


class NavigatorDeck(unittest.TestCase):
    """G-04 на подменном LLM: запрос последним, одна диаграмма долей, ряд январь–июнь одним слайдом."""

    def setUp(self):
        self._model = settings.openai_model
        settings.openai_model = "gpt-6-luna"
        llm_client._SCHEMA_REJECTED.clear()

    def tearDown(self):
        settings.openai_model = self._model

    def run_deck(self, fake):
        req = DeckRequest(input={"topic": "Итоги пилота «Навигатор»",
                                 "material": {"kind": "text", "ref": "redis:x"}},
                          source_mode="strict", audience="management", slides_count=9)
        with mock.patch.object(pipeline, "pptx_to_pdf", side_effect=_pdf):
            return asyncio.run(pipeline.generate_deck("dk_01J0000000000000000000000A", req,
                                                      client=LLMClient(fake), material=NAV))

    def test_g04_structure(self):
        fake = FakeOpenAI(outline=NAV_PLAN, content=nav_content)
        res = self.run_deck(fake)
        spec = res.spec
        kinds = [s.kind for s in spec.slides]
        self.assertEqual(kinds[0], "title")
        self.assertEqual(kinds[-1], "closing")
        ask = spec.slides[-2]
        self.assertEqual((ask.role, ask.kind), ("ask", "statement"))
        self.assertIn("4,5 млн", ask.title)
        series = [s for s in spec.slides if s.kind == "chart_series"]
        shares = [s for s in spec.slides if s.kind == "chart_share"]
        self.assertEqual(len(series), 1)
        self.assertEqual(len(shares), 1)
        self.assertEqual(series[0].variant, "chart_series.line_chart")  # ось — месяцы, 6 точек
        self.assertEqual(series[0].data["series"][0]["values"], [310, 720, 1150, 1540, 1820, 2050])
        self.assertEqual(shares[0].data["series"][0]["values"], [46, 27, 18, 9])
        self.assertEqual([x["label"] for x in shares[0].content["legend"]],
                         ["документы", "переписки", "люди", "задачи"])
        self.assertEqual(series[0].footnote, "по данным: ваш текст")
        # «в 6,6 раза» не из источника → перегенерация с ошибками → исправлено
        metrics = next(s for s in spec.slides if s.kind == "metrics")
        self.assertEqual([i["value"] for i in metrics.content["items"]], ["2 400", "8 000"])
        self.assertEqual(metrics.content["items"][1]["tag"], "цель")
        self.assertTrue(any(d.reason == "facts_fixed" for d in spec.degradations))
        self.assertIn("fit", res.usage.by_stage)
        # нативные диаграммы в PPTX
        prs = Presentation(io.BytesIO(res.pptx))
        charts = [sh for sl in prs.slides for sh in sl.shapes if sh.has_chart]
        self.assertEqual(len(charts), 2)
        line = next(c for c in charts if c.name == "chart:line").chart
        self.assertEqual(list(line.plots[0].series[0].values), [310, 720, 1150, 1540, 1820, 2050])
        donut = next(c for c in charts if c.name == "chart:donut").chart
        self.assertFalse(donut.has_title)
        self.assertIn('val="60"', donut._chartSpace.xml.split("holeSize")[1][:20])
        # golden-проверка G-04 по DeckSpec (tests/golden/check_case.py)
        import sys
        sys.path.insert(0, str(_helpers.ROOT / "tests" / "golden"))
        import check_case
        report = check_case.check(spec.model_dump(mode="json"), "G-04")
        self.assertEqual(report.violations, [])
        bad = spec.model_dump(mode="json")
        bad["slides"][-2], bad["slides"][-3] = bad["slides"][-3], bad["slides"][-2]   # запрос не последний
        self.assertTrue(any("последний содержательный" in v for v in check_case.check(bad, "G-04").violations))

    def test_ask_added_when_model_forgets(self):
        plan = fake_outline(n=7, genre="report", kinds=["statement", "metrics", "chart_series", "chart_share",
                                                         "comparison", "process", "conclusion"],
                            asks=[{"text": "Утвердить бюджет на масштабирование — 4,5 млн руб.", "refs": []}],
                            datasets={2: {"id": "ds1", "label_column": "c1", "value_columns": ["c2"], "rows": None},
                                      3: {"id": "ds2", "label_column": "c1", "value_columns": ["c2"], "rows": None}})
        res = self.run_deck(FakeOpenAI(outline=plan, content=nav_content))
        self.assertEqual(res.spec.slides[-2].role, "ask")
        self.assertEqual(len(res.spec.slides), 9)
        self.assertIn("ask_added_by_code", [d.reason for d in res.spec.degradations])

    def test_strict_short_material_warns(self):
        plan = fake_outline(n=4, genre="report", kinds=["statement", "process", "comparison", "conclusion"])
        res = self.run_deck(FakeOpenAI(outline=plan, content=nav_content))
        self.assertEqual(len(res.spec.slides), 6)
        self.assertIn("slides_short:6/9", res.warnings)

    def test_scan_is_deck_error(self):
        req = DeckRequest(input={"topic": "Тема", "material": {"kind": "document", "ref": "x",
                                                            "mime": "application/pdf", "name": "scan.pdf"}})
        pdf = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
               b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF")
        with self.assertRaises(pipeline.DeckError) as ctx:
            asyncio.run(pipeline.generate_deck("dk_01J0000000000000000000000A", req, client=LLMClient(FakeOpenAI()),
                                               material=pdf))
        self.assertEqual(ctx.exception.code, "SCAN_WITHOUT_TEXT")


if __name__ == "__main__":
    unittest.main()
