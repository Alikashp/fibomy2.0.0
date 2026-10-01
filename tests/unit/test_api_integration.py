"""API end-to-end без сети: POST /v1/decks → задача воркера generate_deck_job (настоящий
core.generate_deck на подменном LLM) → GET статус → скачивание PPTX и PDF; webhook с
подписью; ошибки материала в статусе. Задачу из подменной очереди выполняет тест."""
import hashlib
import hmac
import io
import json
import shutil
import unittest
from types import SimpleNamespace
from unittest import mock

import _helpers  # noqa: F401
from _helpers import FakeOpenAI, fake_outline
from _api import ApiHarness

import httpx
from pptx import Presentation

from api import store, webhook
from config import settings
from core import pipeline
from core.llm import client as llm_client
from core.llm.client import LLMClient

TOPIC = "Как работает фотосинтез"


async def fake_pdf(pptx: bytes, **kw) -> bytes:
    return b"%PDF-1.4 fake"


class EndToEnd(unittest.TestCase):

    def setUp(self):
        self._model = settings.openai_model
        settings.openai_model = "gpt-6-luna"
        llm_client._SCHEMA_REJECTED.clear()
        self.h = ApiHarness()
        self.client, self.key = self.h.new_client()
        self.headers = self.h.auth(self.key)
        self.fake = FakeOpenAI()

    def tearDown(self):
        self.h.close()
        settings.openai_model = self._model

    def run_queue(self, real_pdf: bool = False) -> list[dict]:
        """Выполняет задачи из подменной очереди, как воркер."""
        import worker

        async def gen(deck_id, req, progress, material=None):
            return await pipeline.generate_deck(deck_id, req, progress, client=LLMClient(self.fake), material=material)

        out = []
        patches = [mock.patch.object(worker, "generate_deck", side_effect=gen)]
        if not real_pdf:
            patches.append(mock.patch.object(pipeline, "pptx_to_pdf", side_effect=fake_pdf))
        for p in patches:
            p.start()
        try:
            while self.h.pool.jobs:
                name, args, kwargs = self.h.pool.jobs.pop(0)
                ctx = {"bot": SimpleNamespace(), "redis": self.h.pool}
                kwargs = {k: v for k, v in kwargs.items() if not k.startswith("_")}
                fn = {"generate_deck_job": worker.generate_deck_job, "deliver_webhook_job": webhook.deliver_webhook_job}
                out.append({"name": name, "result": self.h.run(fn[name], ctx, *args, **kwargs)})
        finally:
            for p in patches:
                p.stop()
        return out

    def status(self, deck_id: str) -> dict:
        r = self.h.http.get(f"/v1/decks/{deck_id}", headers=self.headers)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def test_topic_deck_end_to_end(self):
        r = self.h.http.post("/v1/decks", headers=self.headers, json={"input": {"topic": TOPIC}, "slides_count": 9})
        deck_id = r.json()["id"]
        self.assertEqual(self.status(deck_id)["status"], "queued")
        results = self.run_queue()
        self.assertEqual(results[0]["result"]["status"], "ok")
        body = self.status(deck_id)
        self.assertEqual(body["status"], "done")
        self.assertEqual(body["slides"], 9)
        self.assertIsNotNone(body["started_at"])
        self.assertIsNotNone(body["finished_at"])
        self.assertIsNone(body["error"])
        pptx = self.h.http.get(f"/v1/decks/{deck_id}/files/pptx", headers=self.headers)
        self.assertEqual(pptx.status_code, 200)
        prs = Presentation(io.BytesIO(pptx.content))
        self.assertEqual(len(prs.slides), 9)
        titles = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
        self.assertIn(TOPIC, titles)
        pdf = self.h.http.get(f"/v1/decks/{deck_id}/files/pdf", headers=self.headers)
        self.assertEqual(pdf.content[:4], b"%PDF")
        # стоимость колоды — по клиенту
        report = self.h.run(store.usage_report, store.day_start(__import__("datetime").datetime.now(
            __import__("datetime").timezone.utc)))
        self.assertEqual(report[0]["client_id"], self.client.id)
        self.assertGreater(report[0]["cost_rub"], 0)
        deck = self.h.run(_load, deck_id)
        self.assertTrue(deck.counted)
        self.assertIn("store_files", deck.durations_ms)

    def test_text_material_short_warning(self):
        outline = fake_outline(n=3, genre="report", kinds=["statement", "bullets", "conclusion"])
        self.fake = FakeOpenAI(outline=outline)
        text = ("Отдел закупок перешёл на электронные заявки. Заявки обрабатываются за один день вместо трёх. "
                "Сотрудники довольны новым порядком работы. ") * 3
        r = self.h.http.post("/v1/decks", headers=self.headers,
                             json={"input": {"topic": "Электронные заявки", "text": text}, "slides_count": 9})
        deck_id = r.json()["id"]
        self.run_queue()
        body = self.status(deck_id)
        self.assertEqual(body["status"], "done", body)
        self.assertEqual(body["slides"], 5)
        self.assertEqual([w["code"] for w in body["warnings"]], ["SLIDES_SHORT"])
        # материал удалён после обработки (ТЗ 6.5)
        material = self.h.run(_load, deck_id).request["input"]["material"]
        from generation.uploads import UploadNotFound, get_upload
        with self.assertRaises(UploadNotFound):
            self.h.run(get_upload, material["ref"])

    def test_material_error_in_status(self):
        r = self.h.http.post("/v1/decks", headers=self.headers, data={"params": json.dumps({"input": {"topic": TOPIC}})},
                             files={"file": ("broken.docx", io.BytesIO(b"not a zip"), "application/octet-stream")})
        deck_id = r.json()["id"]
        self.run_queue()
        body = self.status(deck_id)
        self.assertEqual(body["status"], "failed")
        self.assertEqual(body["error"]["code"], "BAD_FILE")
        self.assertIsNone(body["files"])
        r = self.h.http.get(f"/v1/decks/{deck_id}/files/pptx", headers=self.headers)
        self.assertEqual(r.json()["error"]["code"], "DECK_FAILED")
        # упавшая колода не списывает суточный лимит
        self.assertEqual(self.h.http.get("/v1/usage", headers=self.headers).json()["used_today"], 0)

    def test_storage_failure(self):
        from core.storage import files as deck_files
        r = self.h.http.post("/v1/decks", headers=self.headers, json={"input": {"topic": TOPIC}})
        deck_id = r.json()["id"]
        with mock.patch.object(deck_files, "put_deck_files", side_effect=ConnectionError("redis down")):
            self.run_queue()
        body = self.status(deck_id)
        self.assertEqual((body["status"], body["error"]["code"]), ("failed", "STORAGE_FAILED"))

    @unittest.skipUnless(shutil.which("soffice"), "нужен LibreOffice")
    def test_real_pdf(self):
        r = self.h.http.post("/v1/decks", headers=self.headers, json={"input": {"topic": TOPIC}, "slides_count": 5})
        deck_id = r.json()["id"]
        self.run_queue(real_pdf=True)
        body = self.status(deck_id)
        self.assertEqual(body["status"], "done", body)
        self.assertIsNotNone(body["files"]["pdf"])
        pdf = self.h.http.get(f"/v1/decks/{deck_id}/files/pdf", headers=self.headers)
        self.assertEqual(pdf.content[:5], b"%PDF-")


class Webhook(EndToEnd):

    def test_webhook_signed_and_retried(self):
        r = self.h.http.post("/v1/decks", headers=self.headers,
                             json={"input": {"topic": TOPIC}, "webhook_url": "https://bot2.example.com/hook"})
        deck_id = r.json()["id"]
        received, answers = [], [500, 200]

        def handler(request: httpx.Request):
            received.append(request)
            return httpx.Response(answers.pop(0))

        transport = httpx.MockTransport(handler)
        real_send = webhook.send

        async def send(deck_id, attempt=1):
            return await real_send(deck_id, attempt, transport=transport)

        with mock.patch.object(webhook, "send", side_effect=send):
            self.run_queue()
        self.assertEqual(len(received), 2, "первая попытка 500 → повтор")
        req = received[-1]
        self.assertEqual(str(req.url), "https://bot2.example.com/hook")
        self.assertEqual(req.headers["X-Fibonacci-Attempt"], "2")
        client = self.h.run(store.get_client, self.client.id)
        expected = "sha256=" + hmac.new(client.webhook_secret.encode(), req.content, hashlib.sha256).hexdigest()
        self.assertEqual(req.headers["X-Fibonacci-Signature"], expected)
        body = json.loads(req.content)
        self.assertEqual((body["id"], body["status"]), (deck_id, "done"))
        self.assertTrue(body["files"]["pptx"].endswith(f"/v1/decks/{deck_id}/files/pptx"))

    def test_webhook_on_failure_and_retry_limit(self):
        r = self.h.http.post("/v1/decks", headers=self.headers,
                             data={"params": json.dumps({"input": {"topic": TOPIC},
                                                         "webhook_url": "https://bot2.example.com/hook"})},
                             files={"file": ("broken.pdf", io.BytesIO(b"%PDF-broken"), "application/pdf")})
        self.assertEqual(r.status_code, 202)
        calls = []

        def handler(request):
            calls.append(json.loads(request.content))
            return httpx.Response(503)

        transport = httpx.MockTransport(handler)
        real_send = webhook.send

        async def send(deck_id, attempt=1):
            return await real_send(deck_id, attempt, transport=transport)

        with mock.patch.object(webhook, "send", side_effect=send):
            results = self.run_queue()
        self.assertEqual(len(calls), 1 + len(webhook.RETRY_DELAYS), "первая попытка и 3 повтора")
        self.assertEqual(calls[0]["status"], "failed")
        webhook_jobs = [x for x in results if x["name"] == "deliver_webhook_job"]
        self.assertEqual([x["result"]["attempt"] for x in webhook_jobs], [1, 2, 3, 4])

    def test_no_webhook_without_url(self):
        self.h.http.post("/v1/decks", headers=self.headers, json={"input": {"topic": TOPIC}})
        results = self.run_queue()
        self.assertEqual([x["name"] for x in results], ["generate_deck_job"])


async def _load(deck_id):
    from core.storage import decks as deck_store
    return await deck_store.load(deck_id)


if __name__ == "__main__":
    unittest.main()
