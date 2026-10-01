"""Новый движок (src/core) на реальной модели — сессия 1 (docs/design/09_PLAN.md).

    python tests/golden/run_new_engine.py --out <папка> [--probe-only]

1. Пробный вызов structured outputs (вопрос 6): принимает ли провайдер
   response_format json_schema strict и соблюдает ли модель maxLength / maxItems.
2. Колоды по теме: G-05 (тема на титуле дословно, ТЗ 7.4) и «Как работает
   фотосинтез» (ручная проверка владельца). Проверки: 9 слайдов, титул = тема,
   нет диаграмм, сноска у слайдов с числами, PPTX читается, PDF собрался.

В --out: probe.json, <кейс>.pptx / .pdf / .deckspec.json, new_engine.md (для
комментария в PR). Код возврата: 0 — колоды собраны, 1 — хотя бы одна упала.
Полный golden на новом движке (G-01…G-05 + питч-дек) — сессия 2.
"""
import argparse
import asyncio
import io
import json
import os
import re
import sys
import time
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
os.environ.setdefault("TELEGRAM_BOT_TOKEN", "123456:golden-dummy")

CASES = [
    ("G-05", dict(input={"topic": "Управление требованиями стейкхолдеров в ИТ-стартапе"},
                  audience="students", slides_count=9, watermark=True)),
    ("photosynthesis", dict(input={"topic": "Как работает фотосинтез"}, audience="general", slides_count=9,
                            watermark=True, author={"name": "Иван Петров", "group": "ББИ-21"})),
]
FOOTNOTE = "Оценочные данные — проверьте перед показом"


async def probe() -> dict:
    """json_schema strict: принят ли и соблюдены ли пределы схемы."""
    from openai import APIStatusError

    from core.llm.client import LLMClient, stage_model
    from core.llm.models import completion_params

    model, effort = stage_model("content")
    schema = {"type": "object", "additionalProperties": False, "required": ["items"], "properties": {
        "items": {"type": "array", "minItems": 1, "maxItems": 3, "items": {
            "type": "object", "additionalProperties": False, "required": ["text"],
            "properties": {"text": {"type": "string", "maxLength": 40}}}}}}
    params = completion_params(800, model=model, effort=effort)
    params["response_format"] = {"type": "json_schema", "json_schema": {"name": "probe", "strict": True, "schema": schema}}
    messages = [{"role": "user", "content": "Перечисли 6 фактов о городском транспорте, каждый — развёрнутым "
                                            "предложением не короче 120 знаков. Ответ — JSON {\"items\": [{\"text\"}]}."}]
    client = LLMClient()._client
    result = {"model": model, "effort": effort}
    started = time.perf_counter()
    try:
        resp = await client.chat.completions.create(messages=messages, **params)
    except APIStatusError as e:
        result.update(accepted=False, status=e.status_code, error=str(e)[:500])
        return result
    result["seconds"] = round(time.perf_counter() - started, 1)
    content = resp.choices[0].message.content or ""
    result.update(accepted=True, finish_reason=resp.choices[0].finish_reason, raw=content[:2000])
    try:
        items = json.loads(content)["items"]
        result.update(valid_json=True, items=len(items), max_items_ok=len(items) <= 3,
                      max_len=max((len(i["text"]) for i in items), default=0),
                      max_length_ok=all(len(i["text"]) <= 40 for i in items))
    except Exception as e:
        result.update(valid_json=False, error=str(e)[:300])
    return result


def check(spec, pptx: bytes, pdf: bytes | None, topic: str) -> list[str]:
    from pptx import Presentation
    problems = []
    if len(spec.slides) != 9:
        problems.append(f"слайдов {len(spec.slides)}, ожидалось 9")
    if " ".join(spec.slides[0].title.split()) != " ".join(topic.split()):
        problems.append(f"титул не равен теме: «{spec.slides[0].title}»")
    if any(s.kind.startswith("chart_") for s in spec.slides):
        problems.append("диаграмма в режиме «по теме»")
    for s in spec.slides[1:-1]:
        has_digit = re.search(r"\d", json.dumps(s.content, ensure_ascii=False) + s.title)
        if has_digit and s.footnote != FOOTNOTE:
            problems.append(f"{s.id}: число без сноски об оценке")
    prs = Presentation(io.BytesIO(pptx))
    if len(prs.slides) != len(spec.slides):
        problems.append("число слайдов PPTX не равно DeckSpec")
    if pdf is None:
        problems.append("PDF не собрался")
    elif len(re.findall(rb"/Type\s*/Page(?!s)", pdf)) != len(spec.slides):
        problems.append("число страниц PDF не равно числу слайдов")
    return problems


async def run_case(name: str, params: dict, out: Path) -> dict:
    from core.models.ids import new_deck_id
    from core.models.request import DeckRequest
    from core.pipeline import generate_deck

    req = DeckRequest.model_validate(params)
    deck_id = new_deck_id()
    started = time.perf_counter()
    try:
        res = await generate_deck(deck_id, req)
    except Exception as e:
        traceback.print_exc()
        return {"case": name, "ok": False, "error": f"{type(e).__name__}: {e}"[:300]}
    seconds = round(time.perf_counter() - started, 1)
    (out / f"{name}.pptx").write_bytes(res.pptx)
    if res.pdf:
        (out / f"{name}.pdf").write_bytes(res.pdf)
    (out / f"{name}.deckspec.json").write_text(res.spec.model_dump_json(indent=2), encoding="utf-8")
    problems = check(res.spec, res.pptx, res.pdf, req.input.topic)
    print(f"\n=== {name}: {seconds} с, {res.usage.cost_rub} ₽, проблем: {len(problems)}")
    for s in res.spec.slides:
        print(f"  {s.index:>2} {s.variant:<32} {s.title}")
        for key, value in s.content.items():
            print(f"       {key}: {value}")
    return {"case": name, "ok": True, "seconds": seconds, "cost_rub": res.usage.cost_rub,
            "stages": res.durations_ms, "usage": res.usage.as_dict(), "problems": problems,
            "degradations": [d.model_dump() for d in res.spec.degradations],
            "variants": [s.variant for s in res.spec.slides], "title": res.spec.slides[0].title}


def summary(probe_result: dict, results: list[dict]) -> str:
    lines = ["## Новый движок (сессия 1)", "", "### Structured outputs (вопрос 6)", ""]
    p = probe_result
    if not p.get("accepted"):
        lines.append(f"- `json_schema` strict **не принят** провайдером ({p.get('status')}): клиент работает "
                     f"через `json_object` + проверку pydantic. Ошибка: `{p.get('error', '')[:200]}`")
    else:
        lines.append(f"- `json_schema` strict **принят** ({p['model']}, effort `{p['effort']}`, {p.get('seconds')} с).")
        if p.get("valid_json"):
            lines.append(f"- `maxItems: 3` соблюдён: **{'да' if p['max_items_ok'] else 'нет'}** (элементов {p['items']}); "
                         f"`maxLength: 40` соблюдён: **{'да' if p['max_length_ok'] else 'нет'}** (самый длинный — {p['max_len']}).")
        else:
            lines.append(f"- ответ не JSON: {p.get('error')}")
    lines += ["", "### Колоды по теме", "",
              "| Кейс | Время, с | ₽ | OUTLINE, с | CONTENT, с | CONVERT, с | Проверки | Деградации |",
              "|---|---|---|---|---|---|---|---|"]
    for r in results:
        if not r["ok"]:
            lines.append(f"| {r['case']} | — | — | — | — | — | ❌ {r['error']} | — |")
            continue
        st = r["stages"]
        sec = lambda k: f"{st.get(k, 0) / 1000:.1f}"  # noqa: E731
        checks = "✅" if not r["problems"] else "❌ " + "; ".join(r["problems"])
        lines.append(f"| {r['case']} | {r['seconds']} | {r['cost_rub']} | {sec('outline')} | {sec('content')} | "
                     f"{sec('convert')} | {checks} | {len(r['degradations'])} |")
    costs = [r["cost_rub"] for r in results if r.get("ok") and r.get("cost_rub") is not None]
    if costs:
        avg = sum(costs) / len(costs)
        lines += ["", f"Средняя стоимость колоды: **{avg:.2f} ₽** (бюджет — до 3 ₽, порог сигнала — 4 ₽, вопрос 14)."]
    lines += ["", "Файлы PPTX и PDF — в артефакте `golden-decks`, папка `new_engine/`."]
    return "\n".join(lines) + "\n"


async def main(out: Path, probe_only: bool) -> int:
    out.mkdir(parents=True, exist_ok=True)
    probe_result = await probe()
    (out / "probe.json").write_text(json.dumps(probe_result, ensure_ascii=False, indent=2), encoding="utf-8")
    print("probe:", json.dumps({k: v for k, v in probe_result.items() if k != "raw"}, ensure_ascii=False))
    results = [] if probe_only else [await run_case(name, params, out) for name, params in CASES]
    (out / "results.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "new_engine.md").write_text(summary(probe_result, results), encoding="utf-8")
    return 0 if all(r["ok"] for r in results) else 1


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--probe-only", action="store_true")
    args = parser.parse_args()
    sys.exit(asyncio.run(main(args.out, args.probe_only)))
