"""Реальные генерации golden-корпуса без Telegram (workflow «golden»).

    python tests/golden/run_golden.py --out <папка> [--only G-04,G-05] [--outline-effort medium]

Новый движок (src/core.pipeline.generate_deck): G-01 docx ×2, G-02 pdf ×2, G-03 по теме,
G-04 текст «Навигатора» ×3, G-05 по теме. Питч-дек — старый движок до сессии 3 (D-038).
Каждая колода проверяется tests/golden/check_case.py по DeckSpec и по готовому
PPTX (открывается python-pptx, диаграммы нативные и совпадают с DeckSpec.data,
шрифт Arial, автоподбор выключен, PDF собран, страниц столько же, сколько слайдов).

--outline-effort — уровень рассуждений OUTLINE (LLM_EFFORT_OUTLINE) для сравнения
low / medium (вопрос 13). Метка уровня — в заголовке отчёта и в маркере комментария.

Нужны переменные окружения, как у воркера: OPENAI_API_KEY, OPENAI_BASE_URL,
OPENAI_MODEL, OPENAI_REASONING_EFFORT (см. .env.example). Для PDF — LibreOffice.

В папке --out: <прогон>.deckspec.json, .pptx, .pdf (питч-дек — .json и .pdf старого
движка), results.json, summary.md. Код возврата: 0 — все генерации прошли
(нарушения check_case на код не влияют), 1 — хотя бы одна упала, 2 — API недоступен.
"""
import argparse
import asyncio
import io
import json
import os
import re
import sys
import time
import traceback
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests" / "golden"))
os.environ.setdefault("TELEGRAM_BOT_TOKEN", "123456:golden-dummy")

FIXTURES = ROOT / "tests" / "fixtures"
MIME = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
    ".txt": "text/plain",
}


@dataclass
class Case:
    run_id: str
    case: str | None          # id кейса ТЗ 7.4 для check_case; None — только генерация
    title: str
    request: dict
    file: str | None = None
    repeat: int = 1


@dataclass
class Case:
    run_id: str
    case: str | None          # id кейса ТЗ 7.4 для check_case; None — только генерация
    title: str
    request: dict
    file: str | None = None
    repeat: int = 1
    engine: str = "new"


CASES = [
    Case("G-01", "G-01", "G-01 отчёт о продажах, docx, только мой материал",
         dict(input={"topic": "Итоги продаж за квартал"}, audience="general", slides_count=9, source_mode="strict"),
         file="test_prodazhi_Q3.docx", repeat=2),
    Case("G-02", "G-02", "G-02 ТЗ хакатона, pdf, только мой материал",
         dict(input={"topic": "Техническое задание для хакатона"}, audience="general", slides_count=9,
              source_mode="strict"),
         file="test_hakaton_Q3.pdf", repeat=2),
    Case("G-03", "G-03", "G-03 отчёт о продажах по теме, без файла",
         dict(input={"topic": "Итоги продаж за квартал"}, audience="general", slides_count=9)),
    Case("G-04", "G-04", "G-04 итоги пилота «Навигатор», текст, руководство",
         dict(input={"topic": "Итоги пилота «Навигатор»"}, audience="management", slides_count=9,
              source_mode="strict"),
         file="navigator_pilot.txt", repeat=3),
    Case("G-05", "G-05", "G-05 требования стейкхолдеров по теме, студенты",
         dict(input={"topic": "Управление требованиями стейкхолдеров в ИТ-стартапе"}, audience="students",
              slides_count=9)),
    Case("pitch", None, "Питч-дек по теме (старый движок до сессии 3)",
         dict(topic="Сервис доставки корма для собак по подписке", presentation_type="pitch_deck",
              audience="investors", language="ru"), engine="old"),
]


def _baseline() -> dict:
    return json.loads((ROOT / "tests/golden/baseline.json").read_text(encoding="utf-8"))["violations"]


async def _preflight() -> str | None:
    """Доступен ли API с этой машины. None — да, иначе текст ошибки."""
    import httpx
    from config import settings
    url = settings.openai_base_url.rstrip("/") + "/models"
    try:
        async with httpx.AsyncClient(timeout=20) as client:
            resp = await client.get(url, headers={"Authorization": f"Bearer {settings.openai_api_key}"})
    except Exception as e:  # сеть, DNS, TLS
        return f"{urlparse(url).netloc} недоступен: {type(e).__name__}: {e}"
    if resp.status_code in (401, 403):
        return f"{urlparse(url).netloc} ответил {resp.status_code}: ключ не принят или доступ с этого IP закрыт — {resp.text[:200]}"
    # 404 и прочее: /models у прокси может не быть — это не повод не пробовать генерацию
    print(f"preflight: GET {url} → {resp.status_code}")
    return None


async def _run_new(case: Case, n: int, out: Path) -> dict:
    from core.models.ids import new_deck_id
    from core.models.request import DeckRequest
    from core.pipeline import generate_deck
    import check_case

    run_name = f"{case.run_id}-{n}" if case.repeat > 1 else case.run_id
    result = {"run": run_name, "case": case.case, "title": case.title, "status": "ok", "engine": "new"}
    params = json.loads(json.dumps(case.request))
    material = None
    if case.file:
        path = FIXTURES / case.file
        material = path.read_bytes()
        kind = "text" if path.suffix == ".txt" else "document"
        params["input"]["material"] = {"kind": kind, "ref": f"golden:{case.file}", "mime": MIME[path.suffix],
                                       "name": None if kind == "text" else case.file}
    request = DeckRequest.model_validate(params)
    started = time.perf_counter()
    try:
        res = await generate_deck(new_deck_id(), request, material=material)
    except Exception as e:
        result.update(status="error", error=f"{type(e).__name__}: {e}"[:500],
                      seconds=round(time.perf_counter() - started, 1))
        traceback.print_exc()
        return result
    spec = res.spec.model_dump(mode="json")
    result.update(seconds=round(time.perf_counter() - started, 1), usage=_total_usage(res.usage),
                  usage_by_stage=res.usage.as_dict(), stages=res.durations_ms, slides=len(spec["slides"]),
                  source_genre=spec["meta"].get("genre"), degradations=spec["degradations"],
                  kinds=[s["kind"] for s in spec["slides"]], variants=[s["variant"] for s in spec["slides"]],
                  warnings_deck=res.warnings)
    (out / f"{run_name}.deckspec.json").write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / f"{run_name}.pptx").write_bytes(res.pptx)
    if res.pdf:
        (out / f"{run_name}.pdf").write_bytes(res.pdf)
    _print_deck(run_name, check_case.from_deckspec(spec))
    file_problems = check_files(res.spec, res.pptx, res.pdf)
    if case.case:
        report = check_case.check(spec, case.case)
        result.update(violations=report.violations + file_problems, warnings=report.warnings,
                      passed=len(report.passed))
    else:
        result.update(violations=file_problems, warnings=[], passed=0)
    return result


def _total_usage(deck_usage) -> dict:
    """Сумма токенов и стоимости по этапам — в формате старого отчёта."""
    total = {"input_tokens": 0, "output_tokens": 0, "reasoning_tokens": 0, "cached_input_tokens": 0, "calls": 0}
    for t in deck_usage.by_stage.values():
        for k in total:
            total[k] += getattr(t, k)
    total["cost_rub"] = deck_usage.cost_rub
    return total


def check_files(spec, pptx: bytes, pdf: bytes | None) -> list[str]:
    """Проверки готового файла (07_TESTING.md, 3.2): PPTX читается, диаграммы нативные и равны
    DeckSpec.data, шрифт Arial, автоподбор выключен, нет слайдов-картинок, PDF собран."""
    from pptx import Presentation
    from pptx.oxml.ns import qn
    problems = []
    try:
        prs = Presentation(io.BytesIO(pptx))
    except Exception as e:
        return [f"PPTX не открывается: {e}"]
    if len(prs.slides) != len(spec.slides):
        problems.append(f"в PPTX {len(prs.slides)} слайдов, в DeckSpec {len(spec.slides)}")
    for sl, slide in zip(prs.slides, spec.slides):
        shapes = list(sl.shapes)
        if shapes and all(sh.shape_type == 13 for sh in shapes):
            problems.append(f"слайд {slide.index}: только картинка")
        charts = [sh for sh in shapes if sh.has_chart]
        if slide.data:
            if len(charts) != 1:
                problems.append(f"слайд {slide.index}: нет нативной диаграммы")
            else:
                got = [list(s.values) for s in charts[0].chart.plots[0].series]
                want = [[v for v in s["values"]] for s in slide.data["series"]]
                if [[None if v is None else round(v, 4) for v in s] for s in got] != \
                        [[None if v is None else round(v, 4) for v in s] for s in want]:
                    problems.append(f"слайд {slide.index}: данные диаграммы не равны DeckSpec.data")
        for sh in shapes:
            if not sh.has_text_frame:
                continue
            body = sh.text_frame._txBody.find(qn("a:bodyPr"))
            if body is not None and (body.find(qn("a:normAutofit")) is not None or body.find(qn("a:spAutoFit")) is not None):
                problems.append(f"слайд {slide.index}: автоподбор включён ({sh.name})")
            fonts = {r.font.name for p in sh.text_frame.paragraphs for r in p.runs if r.font.name}
            if fonts - {"Arial"}:
                problems.append(f"слайд {slide.index}: шрифт {fonts - {'Arial'}}")
    if pdf is None:
        problems.append("PDF не собран")
    elif len(re.findall(rb"/Type\s*/Page(?!s)", pdf)) != len(spec.slides):
        problems.append("страниц PDF не столько же, сколько слайдов")
    return problems


async def _run_old(case: Case, n: int, out: Path) -> dict:
    from generation.content_extractor import extract_from_document
    from generation.llm import generate_presentation_structure
    from generation.llm_models import track_usage
    from generation.pdf_renderer import html_to_pdf
    from generation.template_engine import render_presentation
    from schemas.presentation import UserRequest
    import check_case

    run_name = f"{case.run_id}-{n}" if case.repeat > 1 else case.run_id
    result = {"run": run_name, "case": case.case, "title": case.title, "status": "ok", "engine": "old"}
    data = dict(case.request)
    if case.file:
        path = FIXTURES / case.file
        data.update(source_type="document", source_name=case.file,
                    raw_text=await extract_from_document(path.read_bytes(), MIME[path.suffix]))
    request = UserRequest.model_validate(data)

    started = time.perf_counter()
    try:
        with track_usage() as usage:
            presentation = await generate_presentation_structure(request)
    except Exception as e:
        result.update(status="error", error=f"{type(e).__name__}: {e}"[:500],
                      seconds=round(time.perf_counter() - started, 1))
        traceback.print_exc()
        return result
    result["seconds"] = round(time.perf_counter() - started, 1)
    result["usage"] = usage.as_dict()
    result["slides"] = presentation.slide_count
    result["source_genre"] = presentation.meta.source_genre

    spec = presentation.model_dump(mode="json")
    (out / f"{run_name}.json").write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
    _print_deck(run_name, spec)

    try:
        html = render_presentation(presentation, image_urls={}, watermark=False, color_scheme="light")
        has_mermaid = any(s.layout.value == "diagram" for s in presentation.slides)
        pdf = await html_to_pdf(html, has_mermaid=has_mermaid, expected_pages=presentation.slide_count)
        (out / f"{run_name}.pdf").write_bytes(pdf)
    except Exception as e:  # PDF — артефакт для глаз, его ошибка не отменяет проверку колоды
        result["pdf_error"] = f"{type(e).__name__}: {e}"[:300]

    if case.case:
        report = check_case.check(spec, case.case)
        result.update(violations=report.violations, warnings=report.warnings, passed=len(report.passed))
    return result


def _print_deck(run_name: str, spec: dict) -> None:
    """Текст колоды в лог job'а: артефакты бывают недоступны, а лог читается
    через API GitHub — по нему разбираем нарушения check_case."""
    import check_case
    print(f"--- {run_name}: {len(spec['slides'])} слайдов, жанр {spec['meta'].get('source_genre')}")
    for slide in spec["slides"]:
        print(f"  #{slide['index']} [{slide['layout']}] {slide.get('title')!r}")
        for unit in check_case.units_of(slide):
            if unit.kind != "title":
                print(f"      {unit.kind}: {unit.text[:200]}")


def _fmt_tokens(u: dict | None) -> str:
    if not u:
        return "—"
    return f"{u['input_tokens']:,} / {u['output_tokens']:,} ({u['reasoning_tokens']:,})".replace(",", " ")


def _fmt_cost(u: dict | None) -> str:
    if not u or u.get("cost_rub") is None:
        return "н/д"
    return f"{u['cost_rub']:.2f}".replace(".", ",")


def _sec(ms) -> str:
    return f"{(ms or 0) / 1000:.1f}".replace(".", ",")


def summary_md(results: list[dict], meta: dict, preflight_error: str | None = None) -> str:
    baseline = _baseline()
    marker = f"<!-- golden-report:{meta.get('outline_effort') or 'default'} -->"
    lines = [marker, f"## Golden: реальные генерации · OUTLINE effort **{meta.get('outline_effort') or meta['effort']}**",
             "",
             f"Модель **{meta['model']}**, reasoning_effort **{meta['effort']}** (OUTLINE — "
             f"**{meta.get('outline_effort') or meta['effort']}**), API `{meta['api_host']}`, "
             f"промпты `{meta['prompts_version']}`, коммит `{meta['commit'][:7]}`.", ""]
    if preflight_error:
        lines += [f"**API недоступен — генерации не запускались.** {preflight_error}", ""]
        return "\n".join(lines)

    lines += ["| Прогон | Движок | Статус | Слайдов | Нарушений: до → сейчас | Время, с | OUTLINE / CONTENT / FIT, с | "
              "Стоимость, ₽ | Деградаций | Жанр |",
              "|---|---|---|---|---|---|---|---|---|---|"]
    for r in results:
        if r["status"] != "ok":
            lines.append(f"| {r['run']} | {r.get('engine', 'new')} | ❌ ошибка | — | — | {r.get('seconds', '—')} | — | — | — | — |")
            continue
        before = baseline.get(r["case"], "—") if r["case"] else "—"
        viol = f"{before} → **{len(r['violations'])}**" if r["case"] else (
            f"**{len(r['violations'])}** (файл)" if r.get("engine") == "new" else "не проверяется")
        st = r.get("stages") or {}
        stages = f"{_sec(st.get('outline'))} / {_sec(st.get('content'))} / {_sec(st.get('facts'))}" if st else "—"
        degr = len(r.get("degradations") or []) if r.get("engine") == "new" else "—"
        lines.append(f"| {r['run']} | {r.get('engine', 'new')} | ✅ | {r['slides']} | {viol} | {r['seconds']} | {stages} | "
                     f"{_fmt_cost(r.get('usage'))} | {degr} | {r.get('source_genre') or '—'} |")

    ok = [r for r in results if r["status"] == "ok"]
    new = [r for r in ok if r.get("engine") == "new"]
    costs = [r["usage"]["cost_rub"] for r in new if r.get("usage", {}).get("cost_rub") is not None]
    times = sorted(r["seconds"] for r in new)
    avg = sum(costs) / len(costs) if costs else None
    total = f"Итого: {len(ok)} из {len(results)} генераций прошли. Новый движок: {len(new)} колод"
    if avg is not None:
        total += f", средняя стоимость **{avg:.2f} ₽**".replace(".", ",")
    if times:
        p95 = times[min(len(times) - 1, int(round(0.95 * (len(times) - 1))))]
        total += f", время p50 {times[len(times) // 2]} с, p95 {p95} с (без картинок)"
    lines += ["", total + ".", ""]
    if avg is not None and avg > 4:
        lines += ["**Средняя стоимость выше 4 ₽ — порог сигнала владельцу (D-041).**", ""]
    lines += [f"«До» — колоды старого движка (`{json.loads((ROOT / 'tests/golden/baseline.json').read_text())['report']}`); "
              "для G-04 и G-05 базы нет.", ""]

    for r in results:
        if r["status"] != "ok":
            lines += [f"**{r['run']} — ошибка:** `{r['error']}`", ""]
            continue
        if not r["case"] and r.get("engine") != "new":
            continue
        kinds = " → ".join(k for k in r.get("kinds", []))
        lines += [f"<details><summary><b>{r['run']}</b> — нарушений {len(r['violations'])}, "
                  f"предупреждений {len(r['warnings'])}, жанр: {r.get('source_genre') or '—'}</summary>", ""]
        if kinds:
            lines += [f"Слайды: {kinds}", ""]
        lines += [f"- {v}" for v in r["violations"]] or ["- нарушений нет"]
        if r["warnings"]:
            lines += ["", "Предупреждения:"] + [f"- {w}" for w in r["warnings"]]
        if r.get("degradations"):
            lines += ["", "Деградации:"] + [f"- {d['stage']} {d.get('slide_id') or ''}: {d['reason'][:150]}"
                                            for d in r["degradations"]]
        lines += ["", "</details>", ""]
    lines += ["DeckSpec, PPTX и PDF — в артефакте `golden-decks` этого запуска."]
    return "\n".join(lines)


async def main(out: Path, only: set[str] | None, outline_effort: str | None) -> int:
    from config import settings
    from core.llm import prompts as prompts_v2

    if outline_effort:
        settings.llm_effort_outline = outline_effort
    out.mkdir(parents=True, exist_ok=True)
    meta = {"model": settings.openai_model, "effort": settings.openai_reasoning_effort,
            "outline_effort": settings.llm_effort_outline or settings.openai_reasoning_effort,
            "api_host": urlparse(settings.openai_base_url).netloc, "prompts_version": prompts_v2.version(),
            "commit": os.environ.get("GITHUB_SHA", "local")}

    preflight_error = await _preflight()
    if preflight_error:
        (out / "summary.md").write_text(summary_md([], meta, preflight_error), encoding="utf-8")
        print(preflight_error, file=sys.stderr)
        return 2

    results = []
    try:
        for case in CASES:
            if only and case.run_id not in only:
                continue
            for n in range(1, case.repeat + 1):
                print(f"=== {case.title}, прогон {n}", flush=True)
                run = _run_new if case.engine == "new" else _run_old
                results.append(await run(case, n, out))
                r = results[-1]
                print(f"    {r['status']}, {r.get('seconds')} с, нарушений {len(r.get('violations', []))}", flush=True)
    finally:
        if any(c.engine == "old" for c in CASES):
            from generation.pdf_renderer import shutdown_renderer
            await shutdown_renderer()

    (out / "results.json").write_text(json.dumps({"meta": meta, "results": results}, ensure_ascii=False, indent=2),
                                      encoding="utf-8")
    (out / "summary.md").write_text(summary_md(results, meta), encoding="utf-8")
    return 1 if any(r["status"] != "ok" for r in results) else 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--only", help="через запятую: G-01,G-02,G-03,G-04,G-05,pitch")
    parser.add_argument("--outline-effort", help="reasoning_effort для OUTLINE: low, medium… (вопрос 13)")
    args = parser.parse_args()
    sys.exit(asyncio.run(main(args.out, set(args.only.split(",")) if args.only else None, args.outline_effort)))
